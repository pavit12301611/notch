#!/usr/bin/env python3
"""Build the editable Notch design-review PPTX.

Every mockup is drawn with native PowerPoint shapes and text boxes.  The same
scene graph is optionally rendered to PNG with Pillow so the deck can be
proofed on machines without LibreOffice.  No screenshots are embedded.

Usage:
    python3 scripts/build_design_review.py --output Notch-Design-Review.pptx
"""
from __future__ import annotations

import argparse
import math
import os
from pathlib import Path
from dataclasses import dataclass, field

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.ns import qn

try:
    from PIL import Image, ImageDraw, ImageFont
except ImportError:  # pragma: no cover
    Image = ImageDraw = ImageFont = None

SLIDE_W = 13.333333
SLIDE_H = 7.5
DECK_BG = "0B0B0F"
PANEL_BG = "15151B"
PANEL_BG2 = "1C1C24"
STROKE = "2B2B35"
TEXT = "F5F5F7"
MUTED = "A1A1A6"
DIM = "777782"
ACCENT = "30D158"
ACCENT_DIM = "17351F"
WARN = "FFD60A"
ERROR = "FF9F0A"
BLUE = "64D2FF"
PURPLE = "BF5AF2"

# Current v8 notch tokens (source: extension/stylesheet.css).
N_BG = "000000"
N_CARD = "1C1C1E"
N_CONTROL = "2C2C2E"
N_HOVER = "3A3A3C"
N_PRESSED = "48484A"
N_TEXT = "F5F5F7"
N_TEXT2 = "E5E5EA"
N_MUTED = "98989D"
N_MUTED2 = "8E8E93"
N_MUTED3 = "A1A1A6"
N_DISABLED = "636366"
N_ACCENT = "30D158"
N_BORDER = "242426"
N_EMPTY = "0A0A0B"
N_ACTIVE = "292929"

FONT = "Arial"
MONO = "Consolas"
PREVIEW_DPI = 144
FONT_DIR = Path("/usr/share/fonts/truetype/dejavu")


def rgb(hexstr: str) -> RGBColor:
    return RGBColor.from_string(hexstr)


def clamp(v, lo, hi):
    return max(lo, min(hi, v))


@dataclass
class Prim:
    kind: str
    x: float
    y: float
    w: float = 0
    h: float = 0
    fill: str | None = None
    line: str | None = None
    line_w: float = 1
    radius: float = 0
    alpha: float = 1.0
    dash: bool = False
    text: str = ""
    size: float = 12
    color: str = TEXT
    bold: bool = False
    italic: bool = False
    align: str = "left"
    valign: str = "middle"
    wrap: bool = True
    font: str = FONT
    points: list = field(default_factory=list)
    name: str = ""


class Layer:
    """Draws into a PowerPoint shape tree and records primitives for preview."""

    def __init__(self, slide: "Slide", target=None, prefix: str = ""):
        self.slide = slide
        self.target = target if target is not None else slide.pptx.shapes
        self.prefix = prefix
        self.counter = 0

    def _name(self, part: str) -> str:
        self.counter += 1
        base = self.prefix or self.slide.sid
        return f"{base}/{part}/{self.counter:03d}"

    def _record(self, p: Prim):
        self.slide.prims.append(p)

    def rect(self, x, y, w, h, fill=None, line=None, line_w=1, radius=0, alpha=1.0, dash=False, name="rect"):
        p = Prim("rect", x, y, w, h, fill, line, line_w, radius, alpha, dash, name=self._name(name))
        self._record(p)
        shp = self.target.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
            Inches(x), Inches(y), Inches(w), Inches(h),
        )
        shp.name = p.name
        if radius:
            try:
                shp.adjustments[0] = clamp(radius / min(w, h), 0, 0.5)
            except Exception:
                pass
        self._style(shp, fill, line, line_w, alpha, dash)
        return shp

    def ellipse(self, x, y, w, h, fill=None, line=None, line_w=1, alpha=1.0, name="ellipse"):
        p = Prim("ellipse", x, y, w, h, fill, line, line_w, alpha=alpha, name=self._name(name))
        self._record(p)
        shp = self.target.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(w), Inches(h))
        shp.name = p.name
        self._style(shp, fill, line, line_w, alpha, False)
        return shp

    def line(self, x1, y1, x2, y2, color=STROKE, width=1, dash=False, name="line"):
        p = Prim("line", min(x1, x2), min(y1, y2), abs(x2 - x1), abs(y2 - y1),
                 line=color, line_w=width, dash=dash, points=[(x1, y1), (x2, y2)], name=self._name(name))
        self._record(p)
        shp = self.target.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        shp.name = p.name
        shp.line.color.rgb = rgb(color)
        shp.line.width = Pt(width)
        if dash:
            shp.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        return shp

    def poly(self, points, fill=None, line=None, line_w=1, name="poly"):
        xs = [p[0] for p in points]
        ys = [p[1] for p in points]
        p = Prim("poly", min(xs), min(ys), max(xs) - min(xs), max(ys) - min(ys),
                 fill, line, line_w, points=points, name=self._name(name))
        self._record(p)
        scale = 914400
        builder = self.target.build_freeform(Emu(int(points[0][0] * scale)), Emu(int(points[0][1] * scale)), scale=1.0)
        builder.add_line_segments([(Emu(int(x * scale)), Emu(int(y * scale))) for x, y in points[1:]], close=True)
        shp = builder.convert_to_shape()
        shp.name = p.name
        self._style(shp, fill, line, line_w, 1.0, False)
        return shp

    def text(self, x, y, w, h, text, size=12, color=TEXT, bold=False, align="left", valign="middle",
             wrap=True, font=FONT, italic=False, name="text", alpha=1.0):
        p = Prim("text", x, y, w, h, text=str(text), size=size, color=color, bold=bold, italic=italic,
                 align=align, valign=valign, wrap=wrap, font=font, alpha=alpha, name=self._name(name))
        self._record(p)
        box = self.target.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        box.name = p.name
        tf = box.text_frame
        tf.word_wrap = wrap
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}[valign]
        lines = str(text).split("\n")
        for i, ln in enumerate(lines):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
            para.space_before = Pt(0)
            para.space_after = Pt(0)
            para.line_spacing = 1.0
            run = para.add_run()
            run.text = ln
            f = run.font
            f.name = font
            f.size = Pt(size)
            f.bold = bold
            f.italic = italic
            f.color.rgb = rgb(color)
        return box

    @staticmethod
    def _style(shp, fill, line, line_w, alpha, dash):
        if fill:
            shp.fill.solid()
            shp.fill.fore_color.rgb = rgb(fill)
            if alpha < 1:
                srgb = shp.fill.fore_color._xFill.find(qn("a:srgbClr"))
                if srgb is not None:
                    a = srgb.makeelement(qn("a:alpha"), {"val": str(int(alpha * 100000))})
                    srgb.append(a)
        else:
            shp.fill.background()
        if line:
            shp.line.color.rgb = rgb(line)
            shp.line.width = Pt(line_w)
            if dash:
                shp.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        else:
            shp.line.fill.background()


class Slide:
    def __init__(self, deck: "Deck", sid: str, section: str, title: str, subtitle: str = ""):
        self.deck = deck
        self.sid = sid
        self.section = section
        self.title = title
        self.subtitle = subtitle
        self.pptx = deck.prs.slides.add_slide(deck.prs.slide_layouts[6])
        self.prims: list[Prim] = []
        self.notes = ""
        self.root = Layer(self, prefix=sid)
        self._chrome()

    def _chrome(self):
        c = self.root
        c.rect(0, 0, SLIDE_W, SLIDE_H, fill=DECK_BG, name="slide-bg")
        c.rect(0, 0, SLIDE_W, 0.06, fill=ACCENT, name="top-accent")
        c.text(0.55, 0.28, 8, 0.22, self.section.upper(), size=9, color=ACCENT, bold=True, wrap=False, name="section")
        c.text(0.55, 0.55, 11.8, 0.55, self.title, size=27, color=TEXT, bold=True, wrap=False, name="title")
        if self.subtitle:
            c.text(0.55, 1.12, 12.2, 0.34, self.subtitle, size=11.5, color=MUTED, wrap=False, name="subtitle")
        c.line(0.55, 6.98, 12.78, 6.98, color=STROKE, width=0.75, name="footer-rule")
        c.text(0.55, 7.08, 4, 0.22, "NOTCH  /  DESIGN REVIEW  /  v8 BASELINE", size=8, color=DIM, wrap=False, name="footer-left")
        c.text(10.4, 7.08, 1.5, 0.22, self.sid, size=8, color=MUTED, align="right", wrap=False, name="slide-id")
        c.text(12.0, 7.08, 0.78, 0.22, f"{len(self.deck.slides) + 1:02d}", size=8, color=MUTED, align="right", wrap=False, name="slide-no")

    def group(self, name: str) -> Layer:
        g = self.pptx.shapes.add_group_shape()
        g.name = f"{self.sid}/{name}"
        return Layer(self, target=g.shapes, prefix=f"{self.sid}/{name}")

    def set_notes(self, text: str):
        self.notes = text
        self.pptx.notes_slide.notes_text_frame.text = text


class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width = Inches(SLIDE_W)
        self.prs.slide_height = Inches(SLIDE_H)
        cp = self.prs.core_properties
        cp.title = "Notch — Editable Design Review"
        cp.subject = "v8 baseline design handoff for Fedora GNOME notch"
        cp.author = "Notch project"
        cp.keywords = "notch, GNOME, GTK4, Libadwaita, design review, editable"
        cp.comments = "Native editable shapes and text; no screenshots. Slide IDs N01–N44 are handoff keys."
        self.slides: list[Slide] = []
        self.warnings: list[str] = []

    def slide(self, sid, section, title, subtitle="") -> Slide:
        s = Slide(self, sid, section, title, subtitle)
        self.slides.append(s)
        return s

    def save(self, path: Path):
        path.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(path)

    def preview(self, outdir: Path):
        if Image is None:
            raise RuntimeError("Pillow is required for previews")
        outdir.mkdir(parents=True, exist_ok=True)
        for idx, s in enumerate(self.slides, 1):
            img = render_slide(s, self.warnings)
            img.save(outdir / f"slide-{idx:02d}.png")
        # Contact sheets, 12 slides each.
        thumbs = []
        for idx, s in enumerate(self.slides, 1):
            img = Image.open(outdir / f"slide-{idx:02d}.png")
            img.thumbnail((480, 270))
            thumbs.append(img)
        for start in range(0, len(thumbs), 12):
            batch = thumbs[start:start + 12]
            sheet = Image.new("RGB", (480 * 3, 270 * 4), (11, 11, 15))
            for i, im in enumerate(batch):
                sheet.paste(im, ((i % 3) * 480, (i // 3) * 270))
            sheet.save(outdir / f"contact-{start // 12 + 1:02d}.png")


_font_cache = {}


def _font(size_pt, bold=False, italic=False, font=FONT):
    px = max(1, int(round(size_pt * PREVIEW_DPI / 72)))
    if font == MONO:
        fname = "DejaVuSansMono-Bold.ttf" if bold else "DejaVuSansMono.ttf"
    else:
        fname = "DejaVuSans-Bold.ttf" if bold else "DejaVuSans.ttf"
        if italic:
            fname = "DejaVuSans-Oblique.ttf" if not bold else "DejaVuSans-BoldOblique.ttf"
    path = FONT_DIR / fname
    if not path.exists():
        path = FONT_DIR / "DejaVuSans.ttf"
    key = (str(path), px)
    if key not in _font_cache:
        _font_cache[key] = ImageFont.truetype(str(path), px)
    return _font_cache[key]


def _wrap(draw, text, font, max_w, wrap):
    out = []
    for para in str(text).split("\n"):
        if not wrap:
            out.append(para)
            continue
        words = para.split(" ")
        cur = ""
        for word in words:
            trial = (cur + " " + word).strip()
            if draw.textlength(trial, font=font) <= max_w or not cur:
                cur = trial
            else:
                out.append(cur)
                cur = word
        out.append(cur)
    return out


def render_slide(s: Slide, warnings: list[str]):
    W = int(SLIDE_W * PREVIEW_DPI)
    H = int(SLIDE_H * PREVIEW_DPI)
    img = Image.new("RGBA", (W, H), (11, 11, 15, 255))
    for p in s.prims:
        layer = Image.new("RGBA", (W, H), (0, 0, 0, 0))
        d = ImageDraw.Draw(layer)
        x, y, w, h = [int(round(v * PREVIEW_DPI)) for v in (p.x, p.y, p.w, p.h)]
        fill = _rgba(p.fill, p.alpha) if p.fill else None
        line = _rgba(p.line, p.alpha) if p.line else None
        lw = max(1, int(round(p.line_w * PREVIEW_DPI / 72)))
        if p.kind == "rect":
            if p.radius:
                d.rounded_rectangle([x, y, x + w, y + h], radius=int(p.radius * PREVIEW_DPI), fill=fill, outline=line, width=lw)
            else:
                d.rectangle([x, y, x + w, y + h], fill=fill, outline=line, width=lw)
            if p.dash and line:
                _dash_rect(d, x, y, w, h, line, lw, int(p.radius * PREVIEW_DPI))
        elif p.kind == "ellipse":
            d.ellipse([x, y, x + w, y + h], fill=fill, outline=line, width=lw)
        elif p.kind == "line":
            pts = [(int(a * PREVIEW_DPI), int(b * PREVIEW_DPI)) for a, b in p.points]
            if p.dash:
                _dash_line(d, pts[0], pts[1], line, lw)
            else:
                d.line(pts, fill=line, width=lw)
        elif p.kind == "poly":
            pts = [(int(a * PREVIEW_DPI), int(b * PREVIEW_DPI)) for a, b in p.points]
            d.polygon(pts, fill=fill, outline=line)
        elif p.kind == "text":
            f = _font(p.size, p.bold, p.italic, p.font)
            lines = _wrap(d, p.text, f, max(1, w), p.wrap)
            lh = int(p.size * PREVIEW_DPI / 72 * 1.18)
            total = lh * len(lines)
            if total > h + 4 and p.text.strip():
                warnings.append(f"{s.sid} {p.name}: text height {total}px > box {h}px :: {p.text[:60]!r}")
            if p.valign == "top":
                ty = y
            elif p.valign == "bottom":
                ty = y + h - total
            else:
                ty = y + (h - total) // 2
            for ln in lines:
                tw = d.textlength(ln, font=f)
                if tw > w + 4 and p.text.strip():
                    warnings.append(f"{s.sid} {p.name}: text width {tw:.0f}px > box {w}px :: {ln[:60]!r}")
                if p.align == "center":
                    tx = x + (w - tw) / 2
                elif p.align == "right":
                    tx = x + w - tw
                else:
                    tx = x
                d.text((int(tx), ty), ln, font=f, fill=_rgba(p.color, p.alpha))
                ty += lh
        img = Image.alpha_composite(img, layer)
    return img.convert("RGB")


def _rgba(hexstr, alpha=1.0):
    h = hexstr.lstrip("#")
    return (int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16), int(255 * alpha))


def _dash_line(d, a, b, color, width, dash=8, gap=6):
    x1, y1 = a
    x2, y2 = b
    length = math.hypot(x2 - x1, y2 - y1)
    if length == 0:
        return
    dx, dy = (x2 - x1) / length, (y2 - y1) / length
    pos = 0
    while pos < length:
        end = min(pos + dash, length)
        d.line([(x1 + dx * pos, y1 + dy * pos), (x1 + dx * end, y1 + dy * end)], fill=color, width=width)
        pos += dash + gap


def _dash_rect(d, x, y, w, h, color, width, radius):
    # Approximate dashed rounded rectangle for preview only.
    _dash_line(d, (x + radius, y), (x + w - radius, y), color, width)
    _dash_line(d, (x + w, y + radius), (x + w, y + h - radius), color, width)
    _dash_line(d, (x + w - radius, y + h), (x + radius, y + h), color, width)
    _dash_line(d, (x, y + h - radius), (x, y + radius), color, width)


# ---------------------------------------------------------------------------
# Shared presentation helpers
# ---------------------------------------------------------------------------

def px(v, z=1.0):
    return v * z / 96.0


def fpx(v, z=1.0):
    return v * z * 0.75


def sidebar(s: Slide, specs: list[str], edit_prompt: str = "Layout / spacing / colors", y=1.72, h=4.95, your_edits=None):
    c = s.root
    x = 9.25
    w = 3.53
    c.rect(x, y, w, h, fill=PANEL_BG, line=STROKE, radius=0.14, name="spec-panel")
    c.text(x + 0.2, y + 0.16, w - 0.4, 0.22, "CURRENT BASELINE", size=9, color=ACCENT, bold=True, wrap=False)
    ty = y + 0.48
    for spec in specs[:4]:
        c.ellipse(x + 0.22, ty + 0.07, 0.06, 0.06, fill=ACCENT)
        c.text(x + 0.38, ty, w - 0.6, 0.5, spec, size=10.5, color=TEXT, valign="top", wrap=True)
        lines = max(1, math.ceil(len(spec) / 36))
        ty += 0.21 * lines + 0.1
    if your_edits:
        ty += 0.04
        c.text(x + 0.2, ty, w - 0.4, 0.2, "YOUR UPLOADED EDIT", size=8.5, color=BLUE, bold=True, wrap=False)
        ty += 0.26
        for edit in your_edits[:2]:
            c.text(x + 0.2, ty, w - 0.4, 0.44, "• " + edit, size=9.5, color=BLUE, valign="top", wrap=True)
            ty += 0.42
    ey = y + h - 1.62
    c.line(x + 0.2, ey - 0.1, x + w - 0.2, ey - 0.1, color=STROKE)
    c.text(x + 0.2, ey, w - 0.4, 0.22, "REVIEW / EDIT HERE", size=9, color=WARN, bold=True, wrap=False)
    c.rect(x + 0.2, ey + 0.3, w - 0.4, 1.08, fill=DECK_BG, line=WARN, line_w=0.75, radius=0.08, name="edit-box")
    c.text(x + 0.32, ey + 0.4, w - 0.64, 0.86,
           f"{edit_prompt}\n[ write exact value or instruction ]", size=10, color=DIM, valign="top", wrap=True, name="edit-note")


def caption(s: Slide, x, y, w, text, color=MUTED):
    s.root.text(x, y, w, 0.28, text, size=10, color=color, align="center", wrap=False)


def badge(s: Slide, x, y, text, color=ACCENT, w=None):
    w = w or (0.14 + len(text) * 0.075)
    c = s.root
    c.rect(x, y, w, 0.28, fill=DECK_BG, line=color, radius=0.14)
    c.text(x, y, w, 0.28, text, size=8.5, color=color, bold=True, align="center", wrap=False)


def icon(c: Layer, name: str, x, y, size=16, color=N_TEXT2, z=1.0):
    """Small editable symbolic-style icon. Coordinates are inches."""
    u = size / 16.0
    lw = max(0.8, 1.2 * u)
    def L(a, b, d, e, col=color, width=lw):
        c.line(x + a * u, y + b * u, x + d * u, y + e * u, color=col, width=width)
    def R(a, b, w, h, fill=None, line=color, radius=0, width=lw):
        c.rect(x + a * u, y + b * u, w * u, h * u, fill=fill, line=line, line_w=width, radius=radius * u)
    def E(a, b, w, h, fill=None, line=color, width=lw):
        c.ellipse(x + a * u, y + b * u, w * u, h * u, fill=fill, line=line, line_w=width)
    def P(pts, fill=color, line=None):
        c.poly([(x + a * u, y + b * u) for a, b in pts], fill=fill, line=line)
    if name == "home":
        P([(8, 2), (2, 7), (4, 7), (4, 14), (7, 14), (7, 10), (9, 10), (9, 14), (12, 14), (12, 7), (14, 7)], fill=None, line=color)
    elif name == "music":
        L(6, 12, 6, 4); L(6, 4, 13, 3); L(13, 3, 13, 11)
        E(3, 10.5, 3.5, 3, fill=color, line=None); E(10, 9.5, 3.5, 3, fill=color, line=None)
    elif name == "bell":
        P([(8, 2), (5, 4), (4, 8), (3, 11), (13, 11), (12, 8), (11, 4)], fill=None, line=color)
        L(6, 13, 10, 13); E(7, 1, 2, 2, fill=color, line=None)
    elif name == "chart":
        L(2, 14, 14, 14); L(3, 11, 6, 7); L(6, 7, 9, 9); L(9, 9, 13, 3)
    elif name == "folder":
        P([(2, 5), (6, 5), (7.5, 7), (14, 7), (14, 13), (2, 13)], fill=None, line=color)
    elif name == "clipboard":
        R(4, 3, 8, 11, radius=1); R(6, 1.5, 4, 2.5, fill=color, line=None, radius=0.6)
        L(6, 7, 10, 7); L(6, 9.5, 10, 9.5)
    elif name == "gear":
        E(5, 5, 6, 6, fill=None, line=color)
        for ang in range(0, 360, 45):
            a = math.radians(ang)
            L(8 + 5 * math.cos(a), 8 + 5 * math.sin(a), 8 + 7 * math.cos(a), 8 + 7 * math.sin(a))
    elif name == "close":
        L(4, 4, 12, 12); L(12, 4, 4, 12)
    elif name == "play":
        P([(5, 3), (13, 8), (5, 13)], fill=color)
    elif name == "pause":
        R(4.5, 3.5, 2.5, 9, fill=color, line=None, radius=0.5); R(9, 3.5, 2.5, 9, fill=color, line=None, radius=0.5)
    elif name == "prev":
        P([(11, 3), (5, 8), (11, 13)], fill=color); R(3, 3.5, 1.6, 9, fill=color, line=None)
    elif name == "next":
        P([(5, 3), (11, 8), (5, 13)], fill=color); R(11.4, 3.5, 1.6, 9, fill=color, line=None)
    elif name == "raise":
        R(3, 5, 10, 8, radius=1); L(6, 2, 13, 2); L(13, 2, 13, 9)
    elif name == "focus":
        E(3, 3, 10, 10, fill=None, line=color); L(4, 12, 12, 4)
    elif name == "night":
        P([(10, 2), (6, 3), (4, 6), (4, 10), (6, 13), (10, 14), (12, 12), (9, 11), (7, 8), (8, 5)], fill=color)
    elif name == "dark":
        E(3, 3, 10, 10, fill=None, line=color); P([(8, 3), (8, 13), (11, 12), (13, 8), (11, 4)], fill=color)
    elif name == "copy":
        R(5, 5, 8, 9, radius=1); R(3, 2, 8, 9, radius=1)
    elif name == "trash":
        L(3, 5, 13, 5); R(5, 5, 6, 9, radius=1); L(6, 3, 10, 3); L(7, 7, 7, 12); L(9, 7, 9, 12)
    elif name == "plus":
        L(8, 3, 8, 13); L(3, 8, 13, 8)
    elif name == "refresh":
        E(3, 3, 10, 10, fill=None, line=color); P([(12, 2), (14, 6), (10, 6)], fill=color)
    elif name == "volume":
        P([(3, 6), (6, 6), (9, 3), (9, 13), (6, 10), (3, 10)], fill=color)
        L(11, 5, 13, 8); L(13, 8, 11, 11)
    elif name == "mute":
        P([(3, 6), (6, 6), (9, 3), (9, 13), (6, 10), (3, 10)], fill=color); L(11, 6, 14, 10); L(14, 6, 11, 10)
    elif name == "cpu":
        R(4, 4, 8, 8, radius=1); R(6.5, 6.5, 3, 3, fill=color, line=None)
        for i in (6, 8, 10):
            L(i, 2, i, 4); L(i, 12, i, 14); L(2, i, 4, i); L(12, i, 14, i)
    elif name == "ram":
        R(2, 5, 12, 6, radius=1)
        for i in (4, 7, 10):
            R(i, 7, 2, 2, fill=color, line=None)
        L(4, 11, 4, 13); L(8, 11, 8, 13); L(12, 11, 12, 13)
    elif name == "battery":
        R(2, 5, 11, 6, radius=1); R(13, 7, 1.5, 2, fill=color, line=None); R(3.5, 6.5, 6, 3, fill=color, line=None)
    elif name == "network":
        L(8, 3, 8, 13); P([(5, 6), (8, 3), (11, 6)], fill=None, line=color); P([(5, 10), (8, 13), (11, 10)], fill=None, line=color)
    elif name == "file":
        P([(4, 2), (9, 2), (12, 5), (12, 14), (4, 14)], fill=None, line=color); L(9, 2, 9, 5); L(9, 5, 12, 5)
    elif name == "drop":
        L(8, 2, 8, 10); P([(5, 7), (8, 10), (11, 7)], fill=None, line=color); L(3, 13, 13, 13)
    elif name == "check":
        L(3, 8, 6.5, 11.5); L(6.5, 11.5, 13, 4)
    elif name == "lock":
        R(4, 7, 8, 7, radius=1); E(5.5, 2.5, 5, 5, fill=None, line=color)
    elif name == "search":
        E(3, 3, 7, 7, fill=None, line=color); L(8.5, 8.5, 13, 13)
    else:
        E(4, 4, 8, 8, fill=None, line=color)


def ui_button(c, x, y, w, h, text="", ic=None, kind="button", z=1.0, active=False, disabled=False, align="center"):
    fill = N_CONTROL
    color = N_TEXT2
    radius = px(10, z)
    if kind == "tab":
        fill = N_ACTIVE if active else None
        color = "FFFFFF" if active else N_MUTED3
    elif kind == "icon":
        fill = None
        color = N_MUTED3
    elif kind == "play":
        fill = N_TEXT
        color = N_BG
        radius = px(18, z)
    elif kind == "action":
        fill = N_TEXT if active else N_CONTROL
        color = N_BG if active else N_TEXT2
    elif kind == "small":
        pass
    elif kind == "link":
        fill = None
        color = N_MUTED3
    if disabled:
        fill = N_CARD
        color = N_DISABLED
    c.rect(x, y, w, h, fill=fill, radius=radius, name=f"button-{text or ic}")
    fs = fpx(12.6 if kind in ("small", "link") else 14, z)
    if ic and text:
        iw = px(16, z)
        gap = px(6, z)
        total = iw + gap + len(text) * fs / 72 * 0.60
        sx = x + (w - total) / 2 if align == "center" else x + px(10, z)
        sx = max(sx, x + px(4, z))
        icon(c, ic, sx, y + (h - iw) / 2, iw, color, z)
        c.text(sx + iw + gap, y, max(px(20, z), w - (sx - x) - iw - gap - px(2, z)), h, text, size=fs, color=color, bold=kind in ("tab", "action"), align="left", wrap=False)
    elif ic:
        iw = px(16, z)
        icon(c, ic, x + (w - iw) / 2, y + (h - iw) / 2, iw, color, z)
    elif text:
        c.text(x + px(6, z), y, w - px(12, z), h, text, size=fs, color=color, bold=kind in ("tab", "action"), align=align, wrap=False)


def slider(c, x, y, w, value=0.5, z=1.0, disabled=False):
    bar_h = px(4, z)
    handle = px(12, z)
    col = N_DISABLED if disabled else N_TEXT
    c.rect(x, y - bar_h / 2, w, bar_h, fill=N_HOVER, radius=bar_h / 2)
    c.rect(x, y - bar_h / 2, w * value, bar_h, fill=col, radius=bar_h / 2)
    c.ellipse(x + w * value - handle / 2, y - handle / 2, handle, handle, fill=col)


def album_art(c, x, y, size, z=1.0, label="ART"):
    c.rect(x, y, size, size, fill="2A1B4A", radius=px(14, z), name="album-art")
    c.ellipse(x + size * 0.1, y + size * 0.12, size * 0.55, size * 0.55, fill=PURPLE, alpha=0.85)
    c.ellipse(x + size * 0.42, y + size * 0.38, size * 0.5, size * 0.5, fill=BLUE, alpha=0.75)
    c.rect(x + size * 0.18, y + size * 0.62, size * 0.64, size * 0.16, fill=ACCENT, alpha=0.9, radius=px(4, z))
    c.text(x, y, size, size, label, size=fpx(12, z), color="FFFFFF", bold=True, align="center", wrap=False)


def notch_points(x, y, w, h, top, bottom):
    pts = [(x, y), (x + w, y)]
    # Right flare: cubic from (w,0) to (w-top, top)
    r = top
    p0 = (x + w, y)
    p1 = (x + w - r / 3, y)
    p2 = (x + w - r, y + r / 3)
    p3 = (x + w - r, y + r)
    for i in range(1, 9):
        t = i / 8
        mt = 1 - t
        pts.append((mt**3 * p0[0] + 3 * mt**2 * t * p1[0] + 3 * mt * t**2 * p2[0] + t**3 * p3[0],
                    mt**3 * p0[1] + 3 * mt**2 * t * p1[1] + 3 * mt * t**2 * p2[1] + t**3 * p3[1]))
    # Bottom right corner
    br = bottom
    cx, cy = x + w - r - br, y + h - br
    for i in range(1, 9):
        a = math.radians(90 * i / 8)
        pts.append((cx + br * math.cos(a), cy + br * math.sin(a)))
    pts.append((x + r + br, y + h))
    # Bottom left corner
    cx, cy = x + r + br, y + h - br
    for i in range(1, 9):
        a = math.radians(90 + 90 * i / 8)
        pts.append((cx + br * math.cos(a), cy + br * math.sin(a)))
    pts.append((x + r, y + r))
    # Left flare back to top-left
    p0 = (x + r, y + r)
    p1 = (x + r, y + r / 3)
    p2 = (x + r / 3, y)
    p3 = (x, y)
    for i in range(1, 8):
        t = i / 8
        mt = 1 - t
        pts.append((mt**3 * p0[0] + 3 * mt**2 * t * p1[0] + 3 * mt * t**2 * p2[0] + t**3 * p3[0],
                    mt**3 * p0[1] + 3 * mt**2 * t * p1[1] + 3 * mt * t**2 * p2[1] + t**3 * p3[1]))
    return pts


def notch_surface(c, x, y, w, h, mode="open", z=1.0, opacity=1.0):
    top = px(19 if mode == "open" else 6, z)
    bottom = px(24 if mode == "open" else 14, z)
    pts = notch_points(x, y, w, h, top, bottom)
    c.poly(pts, fill=N_BG, name=f"notch-surface-{mode}")
    return top, bottom


def panel_bar(c, x, y, w, h=0.34, show_clock=False):
    c.rect(x, y, w, h, fill="1D1D1F", name="gnome-panel")
    c.text(x + 0.16, y, 1.2, h, "Activities", size=9.5, color=N_TEXT, bold=True, wrap=False)
    if show_clock:
        c.text(x + w / 2 - 1, y, 2, h, "Fri Oct 30  14:32", size=9.5, color=N_TEXT, align="center", wrap=False)
    icon(c, "network", x + w - 1.15, y + h / 2 - 0.07, 0.14, N_TEXT)
    icon(c, "volume", x + w - 0.9, y + h / 2 - 0.07, 0.14, N_TEXT)
    icon(c, "battery", x + w - 0.65, y + h / 2 - 0.07, 0.14, N_TEXT)
    c.text(x + w - 0.42, y, 0.3, h, "⏻", size=10, color=N_TEXT, align="center", wrap=False)


def open_header(c, x, y, w, active="home", mode="active", z=1.0):
    tabs = [("home", "Home", "home"), ("media", "Music", "music"), ("notifications", "Alerts", "bell"),
            ("system", "System", "chart"), ("shelf", "Shelf", "folder"), ("clipboard", "Clips", "clipboard")]
    h = px(44, z)
    gap = px(4, z)
    tx = x
    for tid, label, ic in tabs:
        show_label = mode == "labels" or (mode == "active" and tid == active)
        tw = px(32, z)
        if show_label:
            tw += px(10, z) + len(label) * fpx(14, z) / 72 * 0.78
        ui_button(c, tx, y + (h - px(30, z)) / 2, tw, px(30, z), text=label if show_label else "", ic=ic,
                  kind="tab", z=z, active=tid == active)
        tx += tw + gap
    bw = px(32, z)
    ui_button(c, x + w - bw * 2 - px(6, z), y + (h - bw) / 2, bw, bw, ic="gear", kind="icon", z=z)
    ui_button(c, x + w - bw, y + (h - bw) / 2, bw, bw, ic="close", kind="icon", z=z)
    return h


def center_y(hpx, z=1.0, top=1.78, bottom=6.62):
    return top + max(0, (bottom - top - px(hpx, z)) / 2)


def open_notch(s: Slide, page, x=0.55, y=None, wpx=640, hpx=320, z=1.22, active=None, header=None, status=None):
    """Draw an open notch and return content coordinates in inches."""
    if y is None:
        y = center_y(hpx, z)
    if header is None:
        content_w = wpx - 40
        header = "labels" if content_w >= 760 else ("active" if content_w >= 420 else "icons")
    c = s.group(f"notch-{page}")
    w = px(wpx, z)
    h = px(hpx, z)
    notch_surface(c, x, y, w, h, "open", z)
    gutter = px(20, z)
    panel = px(32, z)
    top_pad = px(8, z)
    cx = x + gutter
    cy = y + panel + top_pad
    cw = w - gutter * 2
    hh = open_header(c, cx, cy, cw, active or page, header, z)
    py = cy + hh
    ph = h - panel - top_pad - px(16, z) - hh
    footer = 0
    if status:
        footer = px(status.get("height", 28), z)
        ph -= footer
    return c, cx, py, cw, ph, (x, y, w, h), (cx, y + h - px(16, z) - footer, cw, footer)


def empty_state(c, x, y, w, h, ic, title, body, z=1.0):
    c.rect(x, y, w, h, fill=N_EMPTY, line=N_BORDER, radius=px(14, z), name="empty-state")
    if h < px(110, z):
        iw = px(22, z)
        icon(c, ic, x + (w - iw) / 2, y + px(6, z), iw, N_MUTED2, z)
        c.text(x + px(10, z), y + px(30, z), w - px(20, z), px(18, z), title, size=fpx(13, z), color=N_TEXT, bold=True, align="center", wrap=False)
        c.text(x + px(14, z), y + px(50, z), w - px(28, z), max(px(14, z), h - px(54, z)), body, size=fpx(11, z), color=N_MUTED, align="center", valign="top", wrap=True)
    else:
        iw = px(32, z)
        icon(c, ic, x + (w - iw) / 2, y + px(16, z), iw, N_MUTED2, z)
        c.text(x + px(12, z), y + px(54, z), w - px(24, z), px(20, z), title, size=fpx(14, z), color=N_TEXT, bold=True, align="center", wrap=False)
        c.text(x + px(18, z), y + px(78, z), w - px(36, z), h - px(86, z), body, size=fpx(12.6, z), color=N_MUTED, align="center", valign="top", wrap=True)


def quick_actions(c, x, y, z=1.0, compact=False, states=(False, False, True)):
    items = [("Focus", "focus", "Do Not Disturb"), ("Night", "night", "Night Light"), ("Dark", "dark", "Dark Mode")]
    gap = px(8, z)
    cx = x
    for i, (short, ic, full) in enumerate(items):
        label = short if compact else full
        w = px(78 if compact else 150, z)
        h = px(32 if compact else 36, z)
        ui_button(c, cx, y, w, h, text=label, ic=ic, kind="action", z=z, active=states[i])
        cx += w + gap
    return cx - x - gap


# ---------------------------------------------------------------------------
# Page mockups (logical pixels, scaled by z)
# ---------------------------------------------------------------------------

def draw_home_idle(c, x, y, w, h, z=1.0, compact=False, states=(False, False, True)):
    if compact:
        c.text(x, y, w, px(40, z), "14:32", size=fpx(34, z), color=N_TEXT, bold=True, align="center", wrap=False)
        c.text(x, y + px(42, z), w, px(18, z), "Friday, October 30", size=fpx(14, z), color=N_MUTED, align="center", wrap=False)
        aw = px(78, z) * 3 + px(8, z) * 2
        quick_actions(c, x + (w - aw) / 2, y + px(68, z), z, True, states)
    else:
        c.text(x, y + px(8, z), w * 0.5, px(44, z), "14:32", size=fpx(34, z), color=N_TEXT, bold=True, wrap=False)
        c.text(x, y + px(54, z), w * 0.5, px(20, z), "Friday, October 30", size=fpx(14, z), color=N_MUTED, wrap=False)
        aw = px(78, z) * 3 + px(8, z) * 2
        quick_actions(c, x + w - aw, y + (h - px(32, z)) / 2, z, True, states)


def draw_home_media(c, x, y, w, h, z=1.0, compact=False):
    # Home media row stays horizontal; compact only reduces art and narrow hides Full player.
    art = px(56 if compact else 72, z)
    album_art(c, x, y + (h - art) / 2, art, z)
    tx = x + art + px(14, z)
    tw = w - art - px(14, z)
    c.text(tx, y + px(8, z), tw, px(22, z), "Midnight City", size=fpx(17 if not compact else 15, z), color=N_TEXT, bold=True, wrap=False)
    c.text(tx, y + px(32, z), tw, px(18, z), "M83", size=fpx(12.6, z), color=N_MUTED, wrap=False)
    controls(c, tx, y + px(56, z), z, not compact)


def controls(c, x, y, z=1.0, full_player=True):
    bw = px(32, z)
    gap = px(8, z)
    ui_button(c, x, y, bw, bw, ic="prev", kind="icon", z=z)
    ui_button(c, x + bw + gap, y, px(44, z), bw, ic="pause", kind="play", z=z)
    ui_button(c, x + (bw + gap) * 2, y, bw, bw, ic="next", kind="icon", z=z)
    if full_player:
        ui_button(c, x + (bw + gap) * 2 + bw + gap, y, px(86, z), bw, text="Full player", kind="link", z=z)


def draw_media(c, x, y, w, h, z=1.0, empty=False, narrow=False, player_right=False, paused=False):
    if narrow:
        art = px(56, z)
        album_art(c, x + (w - art) / 2, y, art, z)
        c.text(x, y + art + px(6, z), w, px(20, z), "Midnight City", size=fpx(14.7, z), color=N_TEXT, bold=True, align="center", wrap=False)
        c.text(x, y + art + px(28, z), w, px(18, z), "M83", size=fpx(12.6, z), color=N_MUTED, align="center", wrap=False)
        c.text(x, y + art + px(46, z), w, px(18, z), "Hurry Up, We're Dreaming", size=fpx(12.6, z), color=N_MUTED, align="center", wrap=False)
        ui_button(c, x + (w - px(170, z)) / 2, y + art + px(68, z), px(170, z), px(26, z), text="Spotify  ·  2 players ›", ic="raise", kind="small", z=z, align="left")
        py = y + art + px(104, z)
        c.text(x, py, px(36, z), px(16, z), "1:24", size=fpx(12.6, z), color=N_MUTED, wrap=False)
        slider(c, x + px(42, z), py + px(8, z), w - px(84, z), 0.35, z)
        c.text(x + w - px(42, z), py, px(36, z), px(16, z), "4:03", size=fpx(12.6, z), color=N_MUTED, align="right", wrap=False)
        c.rect(x + w - px(2, z), y, px(3, z), h, fill=N_CONTROL, radius=px(1.5, z))
        c.rect(x + w - px(4, z), y + px(8, z), px(3, z), px(58, z), fill=N_MUTED, radius=px(1.5, z))
        return
    art = px(56 if False else 72, z)
    if empty:
        c.rect(x, y, art, art, fill=N_CARD, radius=px(14, z))
        icon(c, "music", x + art / 2 - px(16, z), y + art / 2 - px(16, z), px(32, z), N_MUTED2, z)
        title, artist, album, player = "Nothing playing", "Start a song in an MPRIS-compatible player", "", "No media player"
    else:
        album_art(c, x, y, art, z)
        title, artist, album, player = "Midnight City", "M83", "Hurry Up, We're Dreaming", "Spotify  ·  2 players ›"
    tx = x + art + px(14, z)
    tw = w - art - px(14, z)
    if player_right:
        tw -= px(180, z)
    c.text(tx, y, tw, px(20, z), title, size=fpx(14.7, z), color=N_TEXT, bold=True, wrap=False)
    c.text(tx, y + px(22, z), tw, px(18, z), artist, size=fpx(12.6, z), color=N_MUTED, wrap=False)
    c.text(tx, y + px(40, z), tw, px(18, z), album, size=fpx(12.6, z), color=N_MUTED, wrap=False)
    if player_right:
        ui_button(c, x + w - px(170, z), y + px(40, z), px(170, z), px(26, z), text=player, ic="raise", kind="small", z=z, disabled=empty, align="left")
    else:
        ui_button(c, tx, y + px(60, z), px(170, z), px(26, z), text=player, ic="raise", kind="small", z=z, disabled=empty, align="left")
    py = y + art + px(10, z)
    c.text(x, py, px(36, z), px(16, z), "0:00" if empty else "1:24", size=fpx(12.6, z), color=N_MUTED, wrap=False)
    slider(c, x + px(42, z), py + px(8, z), w - px(84, z), 0 if empty else 0.35, z, disabled=empty)
    c.text(x + w - px(36, z), py, px(36, z), px(16, z), "0:00" if empty else "4:03", size=fpx(12.6, z), color=N_MUTED, align="right", wrap=False)
    cy = py + px(22, z)
    cw = px(32, z) * 3 + px(44, z) + px(10, z) * 3
    cx = x + (w - cw) / 2
    ui_button(c, cx, cy, px(32, z), px(32, z), ic="prev", kind="button", z=z, disabled=empty)
    ui_button(c, cx + px(42, z), cy, px(44, z), px(32, z), ic="play" if empty or paused else "pause", kind="play", z=z, disabled=empty)
    ui_button(c, cx + px(96, z), cy, px(32, z), px(32, z), ic="next", kind="button", z=z, disabled=empty)
    ui_button(c, cx + px(138, z), cy, px(32, z), px(32, z), ic="raise", kind="button", z=z, disabled=empty)
    vy = cy + px(38, z)
    for label, val, ic in [("Player volume", 1.0 if not empty else 0.0, None), ("System output", 0.65, "volume")]:
        c.text(x, vy, px(100, z), px(18, z), label, size=fpx(12.6, z), color=N_MUTED, wrap=False)
        slider(c, x + px(108, z), vy + px(9, z), w - px(190, z), val, z, disabled=empty and ic is None)
        if ic:
            ui_button(c, x + w - px(32, z), vy - px(5, z), px(32, z), px(28, z), ic=ic, kind="icon", z=z)
        else:
            c.text(x + w - px(44, z), vy, px(44, z), px(18, z), "100%" if not empty else "—", size=fpx(12.6, z), color=N_MUTED, align="right", wrap=False)
        vy += px(28, z)


def draw_alerts(c, x, y, w, h, z=1.0, populated=True, compact=False, indent=0):
    ix = x + px(indent, z)
    c.text(ix, y, w * 0.55, px(22, z), "2 notifications" if populated else "0 notifications", size=fpx(14.7, z), color=N_TEXT, bold=True, wrap=False)
    ui_button(c, x + w - px(100 if not compact else 32, z), y - px(2, z), px(100 if not compact else 32, z), px(28, z),
              text="" if compact else "Clear all", ic="trash", kind="small", z=z, disabled=not populated)
    c.text(ix, y + px(26, z), w, px(18, z), "Friday, October 30", size=fpx(12.6, z), color=N_MUTED, wrap=False)
    ly = y + px(52, z)
    if not populated:
        empty_state(c, x, ly, w, h - px(52, z), "bell", "No notifications", "New alerts will peek here.", z)
        return
    cards = [
        ("Files", "Download complete", "Fedora-Workstation-44.iso is ready.", ["Open folder", "Dismiss"]),
        ("Calendar", "Design review in 15 min", "Notch UI pass · Meeting room 2", ["Join"]),
    ]
    for app, title, body, actions in cards:
        ch = px(86 if len(actions) > 1 else 74, z)
        c.rect(x, ly, w, ch, fill=N_CARD, line=N_BORDER, radius=px(14, z), name="notification-card")
        icon(c, "file" if app == "Files" else "bell", x + px(10, z), ly + px(10, z), px(22, z), N_TEXT2, z)
        c.text(x + px(42, z), ly + px(8, z), w - px(84, z), px(20, z), title, size=fpx(14.7, z), color=N_TEXT, bold=True, wrap=False)
        ui_button(c, x + w - px(34, z), ly + px(6, z), px(28, z), px(28, z), ic="close", kind="icon", z=z)
        c.text(x + px(10, z), ly + px(32, z), w - px(20, z), px(28, z), body, size=fpx(13.3, z), color="C7C7CC", valign="top", wrap=True)
        ax = x + px(10, z)
        for a in actions:
            aw = px(12 + len(a) * 7, z)
            ui_button(c, ax, ly + ch - px(30, z), aw, px(24, z), text=a, kind="small", z=z)
            ax += aw + px(8, z)
        ly += ch + px(8, z)


def draw_system(c, x, y, w, h, z=1.0, columns=4, loading=False, states=(False, False, True)):
    stats = [("Processor", "cpu", "—" if loading else "18%", "Waiting for sample" if loading else "All logical processors"),
             ("Memory", "ram", "—" if loading else "42%", "Waiting for sample" if loading else "6.7 GB / 16.0 GB"),
             ("Battery", "battery", "—" if loading else "AC", "Waiting for sample" if loading else "No system battery"),
             ("Network", "network", "…" if loading else "↓ 1.2 MB/s", "Waiting for sample" if loading else "↑ 84 kB/s · wlp2s0")]
    gap = px(8, z)
    cw = (w - gap * (columns - 1)) / columns
    ch = px(70, z)
    for i, (title, ic, val, detail) in enumerate(stats):
        row, col = divmod(i, columns)
        cx = x + col * (cw + gap)
        cy = y + row * (ch + gap)
        c.rect(cx, cy, cw, ch, fill=N_CARD, line=N_BORDER, radius=px(14, z), name=f"stat-{title}")
        icon(c, ic, cx + px(12, z), cy + px(10, z), px(16, z), N_MUTED3, z)
        c.text(cx + px(34, z), cy + px(8, z), cw - px(44, z), px(18, z), title, size=fpx(12.6, z), color=N_MUTED, wrap=False)
        c.text(cx + px(12, z), cy + px(28, z), cw - px(24, z), px(24, z), val, size=fpx(18 if len(val) > 7 else 20.3, z), color=N_TEXT, bold=True, wrap=False)
        c.text(cx + px(12, z), cy + px(52, z), cw - px(24, z), px(16, z), detail, size=fpx(11, z), color=N_MUTED, valign="top", wrap=True)
    rows = math.ceil(4 / columns)
    ay = y + rows * ch + (rows - 1) * gap + px(10, z)
    c.text(x, ay, w, px(14, z), "QUICK ACTIONS", size=fpx(11.2, z), color=N_MUTED2, bold=True, wrap=False)
    quick_actions(c, x, ay + px(18, z), z, False, states)
    c.text(x, ay + px(58, z), w, px(30, z), "Sampling runs only while this page is visible. Network uses the default IPv4 interface.",
           size=fpx(11.5, z), color=N_MUTED, valign="top", wrap=True)


def draw_system_scroll(c, x, y, w, h, z=1.0, columns=2):
    stats = [("Processor", "cpu", "18%", "All logical processors"), ("Memory", "ram", "42%", "6.7 GB / 16.0 GB"),
             ("Battery", "battery", "AC", "No system battery"), ("Network", "network", "↓ 1.2 MB/s", "↑ 84 kB/s · wlp2s0")]
    gap = px(8, z)
    cw = (w - gap * (columns - 1)) / columns
    ch = px(70, z)
    for i, (title, ic, val, detail) in enumerate(stats):
        row, col = divmod(i, columns)
        cx, cy = x + col * (cw + gap), y + row * (ch + gap)
        if cy + ch > y + h:
            continue
        c.rect(cx, cy, cw, ch, fill=N_CARD, line=N_BORDER, radius=px(14, z))
        icon(c, ic, cx + px(12, z), cy + px(10, z), px(16, z), N_MUTED3, z)
        c.text(cx + px(34, z), cy + px(8, z), cw - px(44, z), px(18, z), title, size=fpx(12.6, z), color=N_MUTED, wrap=False)
        c.text(cx + px(12, z), cy + px(28, z), cw - px(24, z), px(24, z), val, size=fpx(18 if len(val) > 7 else 20.3, z), color=N_TEXT, bold=True, wrap=False)
        c.text(cx + px(12, z), cy + px(52, z), cw - px(24, z), px(16, z), detail, size=fpx(11, z), color=N_MUTED, valign="top", wrap=True)
    c.rect(x + w - px(4, z), y, px(3, z), h, fill=N_CONTROL, radius=px(1.5, z))
    c.rect(x + w - px(4, z), y + px(6, z), px(3, z), px(46, z), fill=N_MUTED, radius=px(1.5, z))


def draw_shelf(c, x, y, w, h, z=1.0, state="ready", items=None, compact=False):
    items = items or []
    c.text(x, y, w * 0.4, px(22, z), "Quick stash", size=fpx(14.7, z), color=N_TEXT, bold=True, wrap=False)
    widths = [px(92 if not compact else 32, z), px(64 if not compact else 32, z)]
    if state != "ready":
        widths.append(px(76 if not compact else 32, z))
    gap = px(8, z)
    bx = x + w - sum(widths) - gap * (len(widths) - 1)
    ui_button(c, bx, y - px(2, z), widths[0], px(28, z), text="" if compact else "Add files", ic="plus", kind="small", z=z)
    ui_button(c, bx + widths[0] + gap, y - px(2, z), widths[1], px(28, z),
              text="" if compact else "Clear", ic="trash", kind="small", z=z, disabled=not items or state != "ready")
    if state != "ready":
        ui_button(c, bx + widths[0] + widths[1] + gap * 2, y - px(2, z), widths[2], px(28, z),
                  text="" if compact else ("Starting…" if state == "starting" else "Retry"), ic="refresh", kind="small", z=z, disabled=state == "starting")
    status = {"ready": "Drop files near the notch, or add them here. Originals stay in place.",
              "starting": "Starting file shelf helper…",
              "offline": "File shelf helper is offline. Press Retry to start it."}[state]
    c.text(x, y + px(26, z), w, px(24, z), status, size=fpx(12.6, z), color=N_MUTED, valign="top", wrap=True)
    ly = y + px(54, z)
    drop_h = px(32, z)
    list_h = h - px(54, z) - drop_h - px(8, z)
    if not items:
        if list_h >= px(58, z):
            empty_state(c, x, ly, w, list_h, "folder", "Shelf is empty",
                        "Stash a few files for later. Originals stay in place." if state == "ready" else "Files appear here once the helper is connected.", z)
        else:
            c.text(x, ly, w, max(px(18, z), list_h), "Shelf is empty — files appear once the helper is connected.", size=fpx(11.5, z), color=N_MUTED, align="center", valign="middle", wrap=True)
    else:
        visible = 2 if list_h < px(100, z) else 3
        rh = px(28, z)
        for i, name in enumerate(items[:visible]):
            ry = ly + i * (rh + px(4, z))
            c.rect(x, ry, w, rh, fill=N_CARD, radius=px(12, z), name="shelf-row")
            ui_button(c, x + px(6, z), ry + px(3, z), w - px(112, z), px(24, z), text=name, ic="file", kind="button", z=z, align="left")
            for j, ic in enumerate(["copy", "folder", "trash"]):
                ui_button(c, x + w - px(96, z) + j * px(30, z), ry + px(3, z), px(24, z), px(24, z), ic=ic, kind="icon", z=z)
        if len(items) > visible:
            c.text(x, ly + visible * (rh + px(4, z)), w, px(12, z), f"+{len(items) - visible} more · scroll", size=fpx(10.5, z), color=N_MUTED, align="center", wrap=False)
    dy = y + h - drop_h
    c.rect(x, dy, w, drop_h, fill="111112", line=N_PRESSED, line_w=1.2, radius=px(14, z), dash=True, name="drop-hint")
    c.text(x + px(10, z), dy, w - px(20, z), drop_h, "Drop files when the GTK target appears. Only references are saved.",
           size=fpx(12.6, z), color=N_MUTED, align="center", wrap=True)


def draw_clipboard(c, x, y, w, h, z=1.0, enabled=False, items=None, compact=False, narrow=False):
    items = items or []
    title = "Clipboard history · on" if enabled else "Clipboard history · manual only"
    c.text(x, y, w - px(232 if not compact else 76, z), px(22, z), title, size=fpx(11.5 if compact else 13.2, z), color=N_TEXT, bold=True, wrap=False)
    bw = px(112 if not compact else 32, z)
    ui_button(c, x + w - bw * 2 - px(8, z), y - px(2, z), bw, px(28, z), text="" if compact else "Capture now", ic="clipboard", kind="small", z=z)
    ui_button(c, x + w - bw, y - px(2, z), bw, px(28, z), text="" if compact else "Clear history", ic="trash", kind="small", z=z, disabled=not items)
    c.text(x, y + px(28, z), w, px(44 if narrow else 32, z), "RAM only. Cleared on lock/logout. Mixed-format and known password-manager clips are skipped.",
           size=fpx(12.6, z), color=N_MUTED, valign="top", wrap=True)
    ly = y + px(80 if narrow else 66, z)
    if not items:
        empty_state(c, x, ly, w, h - px(66, z), "clipboard", "No clipboard history",
                    "Copy plain text to begin." if enabled else "Enable automatic history in Settings, or capture the current text explicitly.", z)
    else:
        for i, (text, time) in enumerate(items[:3]):
            rh = px(52 if narrow else 36, z)
            ry = ly + i * (rh + px(6, z))
            c.rect(x, ry, w, rh, fill=N_CARD, radius=px(12, z), name="clipboard-row")
            if narrow:
                ui_button(c, x + px(6, z), ry + px(4, z), w - px(12, z), px(26, z), text=(text[:28] + "…") if narrow and len(text) > 28 else text, ic="clipboard", kind="button", z=z, align="left")
                c.text(x + px(12, z), ry + px(30, z), w - px(24, z), px(16, z), time, size=fpx(11.5, z), color=N_MUTED, wrap=False)
            else:
                ui_button(c, x + px(6, z), ry + px(5, z), w - px(80, z), px(28, z), text=text, ic="clipboard", kind="button", z=z, align="left")
                c.text(x + w - px(66, z), ry, px(54, z), rh, time, size=fpx(11.5, z), color=N_MUTED, align="right", wrap=False)


def status_footer(c, x, y, w, h, text, z=1.0, error=False, retry=False):
    color = ERROR if error else WARN
    c.text(x, y + px(8, z), w - px(120 if retry else 40, z), h - px(8, z), text, size=fpx(11.9, z), color=color, valign="top", wrap=True)
    bx = x + w - px(108 if retry else 32, z)
    if retry:
        ui_button(c, bx, y + px(6, z), px(64, z), px(26, z), text="Retry", kind="small", z=z)
        bx += px(72, z)
    ui_button(c, bx, y + px(6, z), px(28, z), px(26, z), ic="close", kind="icon", z=z)


# ---------------------------------------------------------------------------
# Libadwaita settings mockups
# ---------------------------------------------------------------------------
ADW_BG = "242428"
ADW_HEADER = "303034"
ADW_CARD = "303034"
ADW_ROW = "3A3A3E"
ADW_TEXT = "FFFFFF"
ADW_MUTED = "C0BFBC"
ADW_ACCENT = "78AEED"


def settings_window(s: Slide, active_tab, groups, x=0.55, y=1.78, w=8.35, h=4.95, title="Notch", subtitle=None, tabs=None):
    c = s.group(f"settings-{active_tab}")
    c.rect(x, y, w, h, fill=ADW_BG, line="454549", radius=0.12, name="adw-window")
    c.rect(x, y, w, 0.46, fill=ADW_HEADER, radius=0.12, name="adw-header")
    c.rect(x, y + 0.24, w, 0.22, fill=ADW_HEADER)
    c.text(x + 0.2, y, w - 0.4, 0.46, title, size=12, color=ADW_TEXT, bold=True, align="center", wrap=False)
    c.ellipse(x + w - 0.34, y + 0.15, 0.16, 0.16, fill=None, line=ADW_MUTED, line_w=1)
    tabs = tabs or ["Appearance", "Behavior", "Privacy & shelf", "About"]
    tw = 1.55 if len(tabs) == 4 else 1.9
    tx = x + (w - tw * 4 - 0.12 * 3) / 2
    ty = y + 0.58
    for tab in tabs:
        act = tab == active_tab
        c.rect(tx, ty, tw, 0.34, fill=ADW_ROW if act else None, radius=0.08)
        c.text(tx, ty, tw, 0.34, tab, size=9.5, color=ADW_TEXT if act else ADW_MUTED, bold=act, align="center", wrap=False)
        tx += tw + 0.12
    cy = ty + 0.5
    if subtitle:
        c.text(x + 0.35, cy, w - 0.7, 0.24, subtitle, size=9, color=ADW_MUTED, wrap=False)
        cy += 0.3
    for group in groups:
        c.text(x + 0.35, cy, w - 0.7, 0.24, group["title"], size=11, color=ADW_TEXT, bold=True, wrap=False)
        cy += 0.26
        if group.get("description"):
            c.text(x + 0.35, cy, w - 0.7, 0.34, group["description"], size=8.5, color=ADW_MUTED, valign="top", wrap=True)
            cy += 0.3
        rows = group["rows"]
        gh = len(rows) * 0.36 + 0.06
        c.rect(x + 0.3, cy, w - 0.6, gh, fill=ADW_CARD, radius=0.1, name="adw-group")
        ry = cy + 0.03
        for i, row in enumerate(rows):
            if i:
                c.line(x + 0.48, ry, x + w - 0.48, ry, color="454549", width=0.5)
            kind = row.get("kind", "switch")
            c.text(x + 0.48, ry + 0.01, w - 3.05, 0.18, row["title"], size=9.5, color=ADW_TEXT, wrap=False, valign="middle")
            if row.get("subtitle"):
                c.text(x + 0.48, ry + 0.19, w - 3.05, 0.14, row["subtitle"], size=7.5, color=ADW_MUTED, wrap=False, valign="middle")
            rx = x + w - 1.75
            if kind == "switch":
                on = row.get("value", True)
                c.rect(rx + 0.65, ry + 0.07, 0.42, 0.22, fill=ADW_ACCENT if on else "5E5C64", radius=0.11)
                c.ellipse(rx + 0.85 if on else rx + 0.67, ry + 0.09, 0.18, 0.18, fill="FFFFFF")
            elif kind == "spin":
                c.rect(rx, ry + 0.04, 1.1, 0.28, fill=ADW_ROW, radius=0.06)
                c.text(rx, ry + 0.04, 0.75, 0.28, str(row.get("value", "")), size=9, color=ADW_TEXT, align="center", wrap=False)
                c.text(rx + 0.78, ry + 0.04, 0.14, 0.28, "−", size=9, color=ADW_MUTED, align="center", wrap=False)
                c.text(rx + 0.94, ry + 0.04, 0.14, 0.28, "+", size=9, color=ADW_MUTED, align="center", wrap=False)
            elif kind == "combo":
                c.rect(rx - 0.75, ry + 0.04, 2.1, 0.28, fill=ADW_ROW, radius=0.06)
                c.text(rx - 0.67, ry + 0.04, 1.8, 0.28, str(row.get("value", "")), size=8.5, color=ADW_TEXT, wrap=False)
                c.text(rx + 1.15, ry + 0.04, 0.16, 0.28, "⌄", size=9, color=ADW_MUTED, align="center", wrap=False)
            elif kind == "button":
                c.rect(rx + 0.25, ry + 0.04, 0.85, 0.28, fill=ADW_ROW, radius=0.06)
                c.text(rx + 0.25, ry + 0.04, 0.85, 0.28, row.get("value", "Open"), size=9, color=ADW_TEXT, align="center", wrap=False)
            elif kind == "slider":
                c.rect(rx - 0.2, ry + 0.16, 1.0, 0.04, fill="5E5C64", radius=0.02)
                c.rect(rx - 0.2, ry + 0.16, 1.0 * row.get("value", 0.5), 0.04, fill=ADW_ACCENT, radius=0.02)
                c.ellipse(rx - 0.2 + 1.0 * row.get("value", 0.5) - 0.07, ry + 0.11, 0.14, 0.14, fill="FFFFFF")
                c.text(rx + 0.9, ry + 0.04, 0.3, 0.28, row.get("label", ""), size=8.5, color=ADW_TEXT, wrap=False)
            elif kind == "text":
                c.text(rx - 0.75, ry + 0.04, 2.1, 0.28, str(row.get("value", "")), size=8, color=ADW_MUTED, align="right", wrap=False)
            ry += 0.36
        cy += gh + 0.18
    # Scroll affordance
    c.rect(x + w - 0.12, y + 1.1, 0.05, h - 1.3, fill="3A3A3E", radius=0.025)
    c.rect(x + w - 0.12, y + 1.2, 0.05, 1.0, fill="77767B", radius=0.025)
    return c


def row(title, kind="switch", value=True, subtitle=None):
    return {"title": title, "kind": kind, "value": value, "subtitle": subtitle}


# ---------------------------------------------------------------------------
# Review-deck helpers
# ---------------------------------------------------------------------------

def fit_z(wpx, hpx, max_w=8.55, max_h=4.40, cap=1.35):
    return min(max_w * 96 / wpx, max_h * 96 / hpx, cap)


def chip(s, x, y, text, color=ACCENT, w=None):
    w = w or max(0.9, 0.16 + len(text) * 0.071)
    c = s.root
    c.rect(x, y, w, 0.26, fill=DECK_BG, line=color, line_w=0.9, radius=0.13)
    c.text(x, y, w, 0.26, text, size=8, color=color, bold=True, align="center", wrap=False)
    return w


def native_table(s, x, y, w, h, headers, rows, widths=None, size=10, name="review-table"):
    """Add a real editable PowerPoint table and record a proof approximation."""
    c = s.root
    nrows, ncols = len(rows) + 1, len(headers)
    shape = s.pptx.shapes.add_table(nrows, ncols, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = f"{s.sid}/{name}"
    table = shape.table
    if widths:
        total = sum(widths)
        for i, cw in enumerate(widths):
            table.columns[i].width = Inches(w * cw / total)
    rh = h / nrows
    for ri, row in enumerate([headers] + rows):
        for ci, val in enumerate(row):
            cell = table.cell(ri, ci)
            cell.margin_left = Inches(0.08)
            cell.margin_right = Inches(0.08)
            cell.margin_top = Inches(0.03)
            cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(PANEL_BG2 if ri == 0 else (PANEL_BG if ri % 2 else DECK_BG))
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = str(val)
            p.font.name = FONT
            p.font.size = Pt(size if ri else size - 0.5)
            p.font.bold = ri == 0
            p.font.color.rgb = rgb(ACCENT if ri == 0 else (TEXT if ci == 0 else MUTED))
            cw = table.columns[ci].width / 914400
            cx = x + sum(table.columns[k].width / 914400 for k in range(ci))
            # Proof-only primitives: the PPTX keeps one real editable table, not a duplicate overlay.
            s.prims.append(Prim("rect", cx, y + ri * rh, cw, rh,
                                fill=PANEL_BG2 if ri == 0 else (PANEL_BG if ri % 2 else DECK_BG),
                                line=STROKE, line_w=0.5, name=f"{s.sid}/table-cell-{ri}-{ci}"))
            s.prims.append(Prim("text", cx + 0.08, y + ri * rh + 0.03, cw - 0.16, rh - 0.06,
                                text=str(val), size=size if ri else size - 0.5,
                                color=ACCENT if ri == 0 else (TEXT if ci == 0 else MUTED),
                                bold=ri == 0, valign="middle", wrap=True, name=f"{s.sid}/table-text-{ri}-{ci}"))
    return shape


def page_slide(d, sid, section, title, subtitle, page, draw, specs, prompt, trigger="", your_edits=None, wpx=None, hpx=None, **kwargs):
    sizes = {"home": (640, 210), "media": (640, 300), "notifications": (560, 320), "system": (640, 260), "shelf": (640, 260), "clipboard": (580, 300)}
    wpx = wpx or sizes[page][0]
    hpx = hpx or sizes[page][1]
    s = d.slide(sid, section, title, subtitle)
    z = fit_z(wpx, hpx)
    w = px(wpx, z)
    x = 0.55 + (8.55 - w) / 2
    y = center_y(hpx, z, 1.78, 6.18)
    c, cx, py, cw, ph, rect, foot = open_notch(s, page, x=x, y=y, wpx=wpx, hpx=hpx, z=z, active=page, **kwargs)
    draw(c, cx, py, cw, ph, z)
    if trigger:
        s.root.text(x, y + px(hpx, z) + 0.14, w, 0.3, trigger, size=10, color=MUTED, align="center", wrap=True)
    sidebar(s, specs, prompt, your_edits=your_edits)
    return s


def section_band(s, label, items):
    c = s.root
    c.rect(0.55, 1.78, 12.23, 0.5, fill=PANEL_BG2, line=STROKE, radius=0.08)
    c.text(0.75, 1.78, 2.2, 0.5, label, size=11, color=ACCENT, bold=True, wrap=False)
    c.text(3.0, 1.78, 9.5, 0.5, items, size=10.5, color=MUTED, wrap=False)


def draw_closed_variant(s, x, y, w, title, note, clock=True, unread=False, wings=False, playing=False, hover=False, match_panel=True):
    c = s.root
    c.text(x, y, w, 0.24, title, size=11.5, color=TEXT, bold=True, wrap=False)
    py = y + 0.34
    panel_bar(c, x, py, w, 0.36, show_clock=not clock)
    nw = 2.05 + (0.68 if wings else 0) + (0.12 if hover else 0)
    nh = 0.36 + (0.04 if hover else 0) if match_panel else 0.46
    nx = x + (w - nw) / 2
    g = s.group(title.lower().replace(" ", "-"))
    notch_surface(g, nx, py, nw, nh, "closed", 0.9)
    if wings:
        album_art(g, nx + 0.12, py + (nh - 0.26) / 2, 0.26, 0.5, "")
    if clock:
        g.text(nx + (0.42 if wings else 0), py, nw - (0.84 if wings else 0), nh, "14:32", size=11, color=N_TEXT, bold=True, align="center", wrap=False)
    if unread:
        g.ellipse(nx + nw / 2 + 0.42, py + nh / 2 - 0.03, 0.06, 0.06, fill=N_TEXT)
    if wings and playing:
        for i, hh in enumerate([0.08, 0.16, 0.11]):
            g.rect(nx + nw - 0.34 + i * 0.06, py + nh / 2 - hh / 2, 0.035, hh, fill=N_ACCENT, radius=0.015)
    c.text(x, py + 0.48, w, 0.42, note, size=9.5, color=MUTED, valign="top", wrap=True)


def draw_peek_card(s, x, y, w, h, title, body, actions=(), z=0.95, long=False):
    g = s.group(title.lower().replace(" ", "-")[:24])
    notch_surface(g, x, y, w, h, "peek", z)
    icon(g, "bell", x + 0.25, y + 0.48, 0.3, N_TEXT2)
    g.text(x + 0.72, y + 0.43, w - 1.25, 0.24, title, size=12, color=N_TEXT, bold=True, wrap=False)
    g.text(x + 0.72, y + 0.72, w - 1.25, 0.62 if long else 0.3, body, size=9.5, color="C7C7CC", valign="top", wrap=True)
    ui_button(g, x + w - 0.48, y + 0.43, 0.28, 0.28, ic="close", kind="icon", z=0.75)
    ax = x + 0.72
    for a in list(actions) + ["Open"]:
        aw = 0.18 + len(a) * 0.075
        ui_button(g, ax, y + h - 0.48, aw, 0.28, text=a, kind="small", z=0.75)
        ax += aw + 0.1
    return g


def draw_file_chooser(s, x, y, w, h):
    c = s.group("gtk-file-chooser")
    c.rect(x, y, w, h, fill=ADW_BG, line="454549", radius=0.12)
    c.rect(x, y, w, 0.44, fill=ADW_HEADER, radius=0.12)
    c.text(x + 0.2, y, w - 0.4, 0.44, "Add files to Notch", size=12, color=ADW_TEXT, bold=True, align="center", wrap=False)
    c.rect(x + 0.2, y + 0.62, 1.5, h - 0.84, fill="2A2A2E", radius=0.08)
    for i, name in enumerate(["Recent", "Home", "Documents", "Downloads", "Other Locations"]):
        c.text(x + 0.35, y + 0.78 + i * 0.36, 1.25, 0.24, name, size=9.5, color=ADW_MUTED, wrap=False)
    c.rect(x + 1.9, y + 0.62, w - 2.1, 0.34, fill=ADW_ROW, radius=0.06)
    c.text(x + 2.05, y + 0.62, w - 2.4, 0.34, "Downloads", size=9.5, color=ADW_TEXT, wrap=False)
    files = [("Fedora-44.iso", "2.1 GB"), ("design.pptx", "318 KB"), ("notes.md", "4 KB")]
    for i, (name, meta) in enumerate(files):
        ry = y + 1.1 + i * 0.44
        c.rect(x + 1.9, ry, w - 2.1, 0.38, fill=ADW_CARD if i == 1 else None, radius=0.06)
        icon(c, "file", x + 2.02, ry + 0.08, 0.2, ADW_MUTED)
        c.text(x + 2.32, ry, w - 3.05, 0.38, name, size=9.5, color=ADW_TEXT, wrap=False)
        c.text(x + w - 0.95, ry, 0.7, 0.38, meta, size=8.5, color=ADW_MUTED, align="right", wrap=False)
    c.rect(x + w - 1.1, y + h - 0.5, 0.9, 0.32, fill=ADW_ACCENT, radius=0.06)
    c.text(x + w - 1.1, y + h - 0.5, 0.9, 0.32, "Stash", size=10, color="000000", bold=True, align="center", wrap=False)


def motion_frame(s, x, y, w, h, label, mode, tab="home", z=0.5):
    c = s.root
    c.text(x, y, w, 0.22, label, size=10, color=ACCENT, bold=True, align="center", wrap=False)
    g = s.group(f"motion-{label}")
    if mode == "closed":
        notch_surface(g, x + (w - 1.7) / 2, y + 0.35, 1.7, 0.32, "closed", 0.7)
        g.text(x + (w - 1.7) / 2, y + 0.35, 1.7, 0.32, "14:32", size=9, color=N_TEXT, bold=True, align="center", wrap=False)
    elif mode == "peek":
        notch_surface(g, x + (w - 2.4) / 2, y + 0.35, 2.4, 0.9, "peek", 0.7)
        g.text(x + (w - 2.4) / 2 + 0.2, y + 0.55, 2.0, 0.2, "Notification", size=8.5, color=N_TEXT, bold=True, wrap=False)
    else:
        nw, nh = 2.8, 1.35
        nx, ny = x + (w - nw) / 2, y + 0.35
        notch_surface(g, nx, ny, nw, nh, "open", 0.65)
        open_header(g, nx + 0.18, ny + 0.28, nw - 0.36, tab, "icons", 0.45)
        g.text(nx + 0.18, ny + 0.72, nw - 0.36, 0.3, tab.title(), size=10, color=N_TEXT, bold=True, align="center", wrap=False)


# ---------------------------------------------------------------------------
# Deck construction
# ---------------------------------------------------------------------------

def build_deck() -> Deck:
    d = Deck()

    s = d.slide("N01", "Design handoff", "Notch — design review workbook", "v8 baseline + your uploaded edits · Fedora 44 / GNOME Wayland · editable native shapes")
    c = s.root
    c.text(0.55, 1.95, 7.6, 0.45, "MacBook-style top-edge notch", size=21, color=TEXT, bold=True, wrap=False)
    c.text(0.55, 2.5, 7.7, 1.15, "Is deck mein har mode, page aur important state ka editable mockup hai.\nTum direct shapes/text edit karo ya REVIEW boxes mein exact instruction likho.\nNext UI pass isi workbook ko source of truth maane ga.", size=13.5, color=MUTED, valign="top")
    chip(s, 0.55, 3.95, "NOT A DYNAMIC ISLAND", ACCENT, 2.35)
    chip(s, 3.05, 3.95, "NATIVE EDITABLE SHAPES", BLUE, 2.55)
    chip(s, 5.75, 3.95, "YOUR EDITS PRESERVED", WARN, 2.45)
    c.text(0.55, 4.55, 7.6, 0.3, "Canonical file: Notch-Design-Review.pptx · Branch: arena/71fe92f9-notch", size=10, color=DIM, wrap=False)
    g = s.group("cover-notch")
    notch_surface(g, 8.55, 2.0, 4.2, 2.5, "open", 1.0)
    open_header(g, 8.8, 2.42, 3.7, "home", "icons", 0.9)
    draw_home_idle(g, 8.8, 3.05, 3.7, 1.1, 0.75, False)
    s.set_notes("Canonical review workbook. Uploaded main-branch PPT was merged into this branch before regeneration. User edits are tracked on N45 and applied to N09/N10/N12/N14/N15; the uploaded curl command is preserved on N61.")

    s = d.slide("N02", "How to review", "Edit the deck, not the code", "PowerPoint / LibreOffice Impress / Google Slides · keep slide IDs stable")
    c = s.root
    steps = [("1", "Pick the slide ID", "Footer mein Nxx ID hai. Usi ID se change request likho."),
             ("2", "Edit mockup directly", "Group select/ungroup karke shape, text, color, size ya spacing change karo."),
             ("3", "Use REVIEW box", "Exact value likho: 72→80 px, radius 14→16, color #30D158, 2 lines max."),
             ("4", "Mark decision", "KEEP / CHANGE / DISCUSS likho. New state ho to closest slide duplicate karke Nxx-A naam do."),
             ("5", "Return .pptx", "N41 decision log fill karo. Implementation N42 mapping se approved edits apply karega.")]
    y = 1.85
    for num, title, body in steps:
        c.ellipse(0.65, y, 0.42, 0.42, fill=ACCENT_DIM, line=ACCENT)
        c.text(0.65, y, 0.42, 0.42, num, size=13, color=ACCENT, bold=True, align="center", wrap=False)
        c.text(1.25, y - 0.02, 4.0, 0.28, title, size=14, color=TEXT, bold=True, wrap=False)
        c.text(1.25, y + 0.28, 7.3, 0.5, body, size=11, color=MUTED, valign="top")
        y += 0.92
    c.rect(9.0, 1.85, 3.78, 4.75, fill=PANEL_BG, line=STROKE, radius=0.14)
    c.text(9.25, 2.05, 3.3, 0.25, "LEGEND", size=10, color=ACCENT, bold=True, wrap=False)
    for i, (col, lab) in enumerate([(ACCENT, "Current v8 baseline"), (BLUE, "Your uploaded edit"), (WARN, "Review / decision field"), (ERROR, "Error / warning state"), (PURPLE, "Optional / experimental")]):
        c.rect(9.25, 2.45 + i * 0.42, 0.22, 0.22, fill=col, radius=0.04)
        c.text(9.6, 2.43 + i * 0.42, 3.0, 0.26, lab, size=11, color=TEXT, wrap=False)
    c.text(9.25, 4.75, 3.3, 1.5, "Mockups schematic hain, screenshots nahi.\nDimensions logical px hain jab tak scale note na ho.\nGreen accent sirf review chrome hai; app accent nahi.", size=10.5, color=MUTED, valign="top")
    s.set_notes("Review protocol. Blue marks imported user edits. Yellow boxes are editable decision fields.")

    s = d.slide("N45", "Imported edits", "Your main-branch upload is preserved", "Detected meaningful edits from Notch-Design-Review.pptx on main")
    rows = [["N09 / N10", "Home idle clock + date", "Moved left toward the content gutter; applied to idle Home mockups."],
            ["N12", "Music player switcher", "Moved to the right side of the metadata block; applied to full Music mockup."],
            ["N14 / N15", "Alerts count + date", "Indented to align with notification-card padding; applied to Alerts mockups."],
            ["N08 → N61", "curl install command", "Pasted command preserved verbatim on a dedicated appendix slide so it does not cover the peek mockup."]]
    native_table(s, 0.55, 1.95, 12.23, 3.3, ["Slide", "Element", "Preserved intent"], rows, [1.4, 2.6, 6.0], size=11)
    s.root.text(0.55, 5.55, 12.23, 0.8, "Minor font-rounding differences from PowerPoint export were normalized back to the deck token scale. No application UI/behavior code was changed by this workbook.", size=11, color=MUTED, valign="top")
    s.set_notes("Source upload: main commit 4719de7, root Notch-Design-Review.pptx. Meaningful shape moves and the added curl textbox were detected by comparing against the generated v8 deck.")

    s = d.slide("N03", "Navigation", "Mode map and review index", "Teen surface modes × six pages · settings, adaptive rules and handoff")
    c = s.root
    modes = [("CLOSED", "185×32 default", "Clock · unread · media wings", "N06–N07, N46"), ("PEEK", "440×144 default", "Notification preview · actions", "N08, N47"), ("OPEN", "page-sized, max 640×320", "Header + page + status footer", "N09–N25, N48–N54")]
    x = 0.55
    for title, dim, desc, ref in modes:
        c.rect(x, 1.9, 3.95, 1.8, fill=PANEL_BG, line=STROKE, radius=0.14)
        c.text(x + 0.22, 2.08, 3.5, 0.28, title, size=15, color=ACCENT, bold=True, wrap=False)
        c.text(x + 0.22, 2.45, 3.5, 0.24, dim, size=11, color=TEXT, wrap=False)
        c.text(x + 0.22, 2.78, 3.5, 0.45, desc, size=10.5, color=MUTED, valign="top")
        c.text(x + 0.22, 3.3, 3.5, 0.22, ref, size=9, color=DIM, wrap=False)
        x += 4.15
    pages = [("Home", "N10–N11, N48"), ("Music", "N12–N13, N49–N50"), ("Alerts", "N14–N15, N51"), ("System", "N16–N18, N52"), ("Shelf", "N19–N22, N53"), ("Clips", "N23–N24, N54")]
    x = 0.55
    for name, ref in pages:
        c.rect(x, 4.0, 1.95, 0.9, fill=PANEL_BG2, line=STROKE, radius=0.12)
        c.text(x, 4.12, 1.95, 0.28, name, size=12.5, color=TEXT, bold=True, align="center", wrap=False)
        c.text(x, 4.48, 1.95, 0.22, ref, size=8.5, color=MUTED, align="center", wrap=False)
        x += 2.07
    section_band(s, "ALSO REVIEW", "Tokens N04 · Geometry N05 · Adaptive N26–N27, N55 · Motion N28, N56 · Settings N30–N36, N57–N58 · Handoff N41–N44")
    s.root.text(0.55, 6.35, 12.23, 0.3, "Review worksheets: N59–N60 · Uploaded install command: N61", size=10.5, color=MUTED, wrap=False)
    s.set_notes("Source: geometry.js PAGE_SIZES/notchGeometry; view.js mode/page stack.")

    s = d.slide("N04", "Design tokens", "Current v8 visual language", "Yahan token edit karo; page mockups isi language ko follow karte hain")
    c = s.root
    colors = [("Surface", N_BG), ("Card", N_CARD), ("Control", N_CONTROL), ("Hover", N_HOVER), ("Pressed", N_PRESSED), ("Text", N_TEXT), ("Muted", N_MUTED), ("Disabled", N_DISABLED), ("Accent bars", N_ACCENT), ("Status", WARN), ("Error", ERROR)]
    for i, (name, hexv) in enumerate(colors):
        cx, cy = 0.55 + (i % 6) * 2.07, 1.9 + (i // 6) * 1.15
        c.rect(cx, cy, 1.9, 0.55, fill=hexv, line=STROKE, radius=0.08)
        c.text(cx, cy + 0.62, 1.9, 0.2, name, size=10, color=TEXT, bold=True, align="center", wrap=False)
        c.text(cx, cy + 0.84, 1.9, 0.2, f"#{hexv}", size=9, color=MUTED, align="center", wrap=False, font=MONO)
    c.text(0.55, 4.45, 5.9, 0.25, "SPACING & RADII", size=10, color=ACCENT, bold=True, wrap=False)
    for i, sp in enumerate(["Gutters 20 px · header 44 px · page spacing 10 px", "List spacing 8 px · card radius 14 / padding 12", "Row radius 12 / padding 8×10 · button radius 10", "Play radius 18 · slider bar 4 / handle 12"]):
        c.text(0.55, 4.8 + i * 0.32, 6.0, 0.26, sp, size=11, color=TEXT, wrap=False)
    c.text(7.0, 4.45, 5.7, 0.25, "TYPE SCALE (base 10.5 pt)", size=10, color=ACCENT, bold=True, wrap=False)
    for i, (a, b) in enumerate([("Home clock", "2.4 em / 700"), ("Home title", "1.2 em / 700"), ("Page title", "1.05 em / 700"), ("Stat value", "1.45 em / 700"), ("Body / muted", "1.0 / 0.9 em"), ("Section label", "0.8 em / 700")]):
        c.text(7.0, 4.8 + i * 0.3, 2.6, 0.24, a, size=10.5, color=TEXT, wrap=False)
        c.text(9.7, 4.8 + i * 0.3, 3.0, 0.24, b, size=10.5, color=MUTED, wrap=False, font=MONO)
    s.set_notes("Source: extension/stylesheet.css; layout.js LAYOUT_LIMITS.")

    s = d.slide("N05", "Geometry", "Shape, size and placement", "Logical pixels at scale 1.0 · monitor-clamped · top flares + rounded bottom")
    c = s.root
    for i, (name, desc) in enumerate([("panel", "Default · replaces centered clock"), ("top-center", "Below panel + 6 px"), ("floating", "Top margin 12 px")]):
        x = 0.55 + i * 4.15
        c.rect(x, 1.85, 3.95, 2.5, fill=PANEL_BG, line=STROKE, radius=0.14)
        panel_bar(c, x + 0.2, 2.05, 3.55, 0.3, show_clock=name != "panel")
        nw, nx = 1.7, x + (3.95 - 1.7) / 2
        ny = 2.05 + (0 if name == "panel" else (0.36 if name == "top-center" else 0.48))
        notch_surface(c, nx, ny, nw, 0.32, "closed", 0.8)
        c.text(nx, ny, nw, 0.32, "14:32", size=9, color=N_TEXT, bold=True, align="center", wrap=False)
        c.text(x + 0.2, 3.35, 3.55, 0.25, name, size=13, color=ACCENT, bold=True, wrap=False)
        c.text(x + 0.2, 3.68, 3.55, 0.4, desc, size=10.5, color=MUTED, valign="top")
    c.text(0.55, 4.65, 12.2, 0.25, "DEFAULT DIMENSIONS", size=10, color=ACCENT, bold=True, wrap=False)
    dims = [("Closed", "185×32", "+64 wings · +12×4 hover"), ("Peek", "440×144", "panel 32 + 112"), ("Home", "640×210", "idle / media"), ("Music", "640×300", "full player"), ("Alerts", "560×320", "list"), ("System", "640×260", "stats + actions"), ("Shelf", "640×260", "files"), ("Clips", "580×300", "history")]
    for i, (a, b, note) in enumerate(dims):
        cx, cy = 0.55 + (i % 4) * 3.12, 5.0 + (i // 4) * 0.78
        c.rect(cx, cy, 2.95, 0.66, fill=PANEL_BG2, line=STROKE, radius=0.08)
        c.text(cx + 0.12, cy + 0.06, 1.0, 0.2, a, size=10, color=TEXT, bold=True, wrap=False)
        c.text(cx + 1.1, cy + 0.06, 1.7, 0.2, b, size=10, color=ACCENT, wrap=False, font=MONO)
        c.text(cx + 0.12, cy + 0.34, 2.7, 0.22, note, size=8.5, color=MUTED, wrap=False)
    s.set_notes("Source: geometry.js; surface.js. Open radii 19/24, closed 6/14. Max open 640×320 default.")

    s = d.slide("N06", "Closed mode", "Idle clock in the top bar", "placement=panel · replace-panel-clock=true · idle-clock=true")
    c = s.root
    panel_bar(c, 0.55, 2.1, 8.35, 0.42)
    g = s.group("closed-idle")
    notch_surface(g, 3.65, 2.1, 2.15, 0.42, "closed", 1.0)
    g.text(3.65, 2.1, 2.15, 0.42, "14:32", size=13, color=N_TEXT, bold=True, align="center", wrap=False)
    caption(s, 0.55, 2.75, 8.35, "GNOME center clock hidden only while notch is visible on the primary monitor.")
    panel_bar(c, 0.55, 3.55, 8.35, 0.42)
    g = s.group("closed-unread")
    notch_surface(g, 3.65, 3.55, 2.15, 0.42, "closed", 1.0)
    g.text(3.65, 3.55, 1.95, 0.42, "14:32", size=13, color=N_TEXT, bold=True, align="center", wrap=False)
    g.ellipse(5.48, 3.72, 0.07, 0.07, fill=N_TEXT)
    caption(s, 0.55, 4.2, 8.35, "Unread dot: notification count > 0 and idle clock enabled.")
    panel_bar(c, 0.55, 5.0, 8.35, 0.42)
    g = s.group("closed-hover")
    notch_surface(g, 3.59, 5.0, 2.27, 0.46, "closed", 1.0)
    g.text(3.59, 5.0, 2.27, 0.46, "14:32", size=13, color=N_TEXT, bold=True, align="center", wrap=False)
    caption(s, 0.55, 5.65, 8.35, "Hover: +12 px width, +4 px height in panel placement; opens after 250 ms if enabled.")
    sidebar(s, ["Base 185×32; min width 150", "Match panel height by default", "Click / Super+` toggles open", "Clock off → blank closed notch"], "Clock size / dot / hover growth")
    s.set_notes("Source: view.js build/updateClock/updateSummary/syncPanelClock; geometry.js.")

    s = d.slide("N07", "Closed mode", "Media wings", "closed-media=true · wings add 64 px · album art stays full-color")
    c = s.root
    for y, playing, label, note in [(2.15, True, "Playing", "20 px art left · clock center · animated 3-bar visualizer right."), (3.65, False, "Paused", "Art remains; visualizer hidden. Stopped/no player returns to idle width."), (5.15, False, "Wings disabled", "Settings off → no art/visualizer even while playing.")]:
        panel_bar(c, 0.55, y, 8.35, 0.42)
        g = s.group(f"closed-media-{label.lower().replace(' ', '-')}")
        nw = 2.15 if label == "Wings disabled" else 2.9
        nx = 0.55 + (8.35 - nw) / 2
        notch_surface(g, nx, y, nw, 0.42, "closed", 1.0)
        if label != "Wings disabled":
            album_art(g, nx + 0.14, y + 0.06, 0.3, 0.6, "")
        g.text(nx + (0.52 if label != "Wings disabled" else 0), y, nw - (1.04 if label != "Wings disabled" else 0), 0.42, "14:32", size=13, color=N_TEXT, bold=True, align="center", wrap=False)
        if playing:
            for i, hh in enumerate([0.1, 0.2, 0.13]):
                g.rect(nx + nw - 0.46 + i * 0.07, y + 0.21 - hh / 2, 0.04, hh, fill=N_ACCENT, radius=0.02)
        elif label == "Paused":
            icon(g, "pause", nx + nw - 0.42, y + 0.13, 0.16, N_MUTED3)
        caption(s, 0.55, y + 0.65, 8.35, note)
    sidebar(s, ["249×32 with wings", "Art 20 px, radius 4", "Bars 3×3 px, spacing 2", "Visualizer only while closed + visible"], "Wing width / art size / bar color")
    s.set_notes("Source: view.js build/updateSummary/syncVisualizer; stylesheet.css notch-live-*.")

    s = d.slide("N08", "Peek mode", "Notification peek", "auto-open-notifications=true · 5000 ms default · critical urgency falls back to native banner")
    draw_peek_card(s, 0.85, 2.0, 5.5, 1.8, "Download complete", "Fedora-Workstation-44.iso is ready.")
    caption(s, 0.85, 4.0, 5.5, "Basic peek: source icon 32 px · title/body · dismiss · Open.")
    draw_peek_card(s, 0.85, 4.55, 5.5, 1.8, "Design review in 15 min", "Notch UI pass · Meeting room 2", ["Join", "Dismiss"])
    caption(s, 0.85, 6.55, 5.5, "Up to 2 native actions + Open. Clicking the notch opens Alerts.")
    sidebar(s, ["440×144 at panel 32", "Icon 32 px, radius 8", "Body wraps; title ellipsizes", "Hover/press postpones timeout"], "Peek height / duration / action order")
    s.set_notes("Source: view.js showPeek/schedulePeek; notifications.js. Uploaded curl command moved to N61 to avoid covering this mockup.")

    s = d.slide("N46", "Closed mode", "Closed variants and fallbacks", "Clock off · unread · hover · custom collapsed height")
    draw_closed_variant(s, 0.55, 1.95, 4.0, "Clock off", "idle-clock=false: closed notch is blank; click/shortcut still opens.", clock=False)
    draw_closed_variant(s, 4.85, 1.95, 4.0, "Unread alert", "Unread dot only appears with idle clock enabled.", unread=True)
    draw_closed_variant(s, 0.55, 4.15, 4.0, "Hover affordance", "Panel placement grows +12×4 px before opening.", hover=True)
    draw_closed_variant(s, 4.85, 4.15, 4.0, "Custom height", "match-panel-height=false uses collapsed-height (24–64 px).", match_panel=False)
    sidebar(s, ["Min closed width 150 px", "Collapsed height 24–64 px", "Hover growth only in panel mode", "Hidden contexts restore GNOME clock"], "Closed-state spacing / fallback")
    s.set_notes("Source: geometry.js notchGeometry; view.js readOptions/syncPanelClock.")

    s = d.slide("N47", "Peek mode", "Peek edge cases", "Long body · native actions · policy fallback")
    draw_peek_card(s, 0.55, 1.95, 4.1, 2.0, "Long notification body", "Long body wraps across multiple lines instead of clipping. The full text remains available in Alerts.", ["Reply"], long=True)
    draw_peek_card(s, 4.85, 1.95, 4.1, 2.0, "Two native actions", "Accept or decline the incoming event.", ["Accept", "Decline"])
    c = s.root
    c.rect(0.55, 4.35, 8.35, 1.9, fill=PANEL_BG, line=STROKE, radius=0.12)
    c.text(0.8, 4.55, 7.85, 0.25, "When the notch does NOT take the banner", size=12, color=ACCENT, bold=True, wrap=False)
    c.text(0.8, 4.95, 7.85, 1.1, "• Critical urgency → native GNOME banner\n• DND / per-app policy suppresses banner → no peek\n• Fullscreen / overview / lock hidden contexts → native path or hidden\n• External file drag → peek suppressed", size=10.5, color=TEXT, valign="top")
    sidebar(s, ["Max 2 native actions + Open", "Body wraps, title ellipsizes", "Timeout pauses on hover/press", "Critical urgency never intercepted"], "Long-text limit / action priority")
    s.set_notes("Source: notifications.js NotificationBridge; view.js showPeek/notificationActive/dragPosition.")

    s = d.slide("N09", "Open mode", "Open-shell anatomy", "Shared chrome for every page · status footer reserves height only when visible")
    c, cx, py, cw, ph, rect, foot = open_notch(s, "home", wpx=640, hpx=320, z=1.22, active="home")
    draw_home_idle(c, cx, py + px(18, 1.22), cw, ph - px(30, 1.22), 1.22)
    s.root.text(0.55, 6.15, 8.35, 0.25, "Panel 32 → top pad 8 → header 44 → page → optional status → bottom pad 16.", size=10, color=BLUE, align="center", wrap=False)
    sidebar(s, ["Gutters 20 px left/right", "Top pad 8 below panel", "Header 44 px; tabs + settings + close", "Page scroll vertical only"], "Header height / gutters / tab order", your_edits=["Home clock/date aligned left in idle mockups."])
    s.set_notes("Source: layout.js contentLayout; view.js build/layoutPages/applyGeometry. User moved idle clock/date left; preserved.")

    page_slide(d, "N10", "Home page", "Home — idle", "640×210 · no MPRIS player · clock + compact quick actions", "home",
               lambda c, x, y, w, h, z: draw_home_idle(c, x, y + px(12, z), w, h - px(24, z), z),
               ["Home clock 2.4 em / 700", "Date 1 em muted", "Compact actions: Focus / Night / Dark", "Active action = white fill, black text"], "Clock/date alignment / action spacing",
               trigger="Clock uses locale time; date is weekday, month day.", your_edits=["Clock and date moved left toward gutter."])
    d.slides[-1].set_notes("Source: view.js buildHome/updateClock/updateActions. Uploaded edit preserved: clock/date left alignment.")

    page_slide(d, "N11", "Home page", "Home — media playing", "640×210 · replaces idle stack when any MPRIS player exists", "home",
               lambda c, x, y, w, h, z: draw_home_media(c, x, y + px(8, z), w, h - px(16, z), z),
               ["Art 72 px, radius 14, padding 6", "Title 1.2 em; artist muted", "Play button white; icons symbolic", "Full player hidden below 340 px content"], "Art size / control spacing / link label",
               trigger="72 px art · title/artist · previous/play-pause/next · Full player link.")
    d.slides[-1].set_notes("Source: view.js buildHome/updateHome/adaptPages.")

    s = d.slide("N48", "Home page", "Home — compact and narrow", "Effective content width <420 px compact · <340 px narrow")
    c, cx, py, cw, ph, rect, foot = open_notch(s, "home", x=0.7, y=2.15, wpx=400, hpx=210, z=1.05, active="home", header="icons")
    draw_home_idle(c, cx, py, cw, ph, 1.05, compact=True)
    caption(s, 0.7, 4.75, 4.2, "Compact: clock above centered actions.")
    c, cx, py, cw, ph, rect, foot = open_notch(s, "home", x=5.1, y=2.15, wpx=340, hpx=210, z=1.05, active="home", header="icons")
    draw_home_media(c, cx, py, cw, ph, 1.05, compact=True)
    caption(s, 5.1, 4.75, 3.55, "Narrow: no Full player; 56 px art.")
    sidebar(s, ["Compact stacks idle clock/actions", "Narrow hides Full player link", "Home art 72→56 px in compact", "Actions remain centered"], "Compact breakpoint / narrow copy")
    s.set_notes("Source: view.js adaptPages; layout.js compact/narrow thresholds.")

    page_slide(d, "N12", "Music page", "Music — full player", "640×300 · art, metadata, seek, transport, player/system volume", "media",
               lambda c, x, y, w, h, z: draw_media(c, x, y + px(4, z), w, h, z, player_right=True),
               ["Art 72 px; metadata wraps", "Seek bar 4 px / handle 12", "Transport centered; raise player", "Player volume + system output/mute"], "Metadata lines / slider labels / volume rows",
               trigger="Seek disabled when canSeek=false or length=0. Switcher enabled only with multiple players.", your_edits=["Player switcher moved to right side."])
    d.slides[-1].set_notes("Source: view.js buildMedia/updateMedia/updateAudio. Uploaded edit preserved: player switcher right aligned.")

    page_slide(d, "N13", "Music page", "Music — no player / unavailable", "Same shell; disabled controls remain visible and readable", "media",
               lambda c, x, y, w, h, z: draw_media(c, x, y + px(4, z), w, h, z, empty=True),
               ["Disabled color #636366", "Disabled fill #1C1C1E", "System output can still work", "No hidden controls when empty"], "Empty copy / disabled contrast",
               trigger="Fallback symbolic art · 'Nothing playing' · 'No media player' · seek/transport disabled.")
    d.slides[-1].set_notes("Source: view.js buildMedia/updateMedia; stylesheet.css :insensitive.")

    s = d.slide("N49", "Music page", "Music — paused, live stream, muted", "Paused transport · disabled seek · system mute")
    c, cx, py, cw, ph, rect, foot = open_notch(s, "media", x=0.7, y=1.95, wpx=560, hpx=300, z=1.05, active="media")
    draw_media(c, cx, py, cw, ph, 1.05, paused=True)
    caption(s, 0.7, 5.35, 5.85, "Paused: play icon replaces pause; metadata and art remain.")
    c = s.root
    c.rect(0.7, 5.75, 8.0, 0.85, fill=PANEL_BG, line=STROKE, radius=0.1)
    c.text(0.9, 5.9, 7.6, 0.55, "Live stream: length=0 → seek disabled, duration 0:00. Muted output: speaker icon changes to muted; volume slider remains available.", size=10.5, color=TEXT, valign="middle")
    sidebar(s, ["canSeek=false disables seek", "length=0 disables seek", "Mute toggles output icon", "Player volume independent"], "Paused / stream / muted states")
    s.set_notes("Source: view.js updateMedia/updateAudio; media.js canSeek/length; audio.js muted.")

    s = d.slide("N50", "Music page", "Music — narrow layout", "Art stacks above metadata; volume rows stack")
    c, cx, py, cw, ph, rect, foot = open_notch(s, "media", x=1.5, y=1.95, wpx=340, hpx=300, z=1.1, active="media", header="icons")
    draw_media(c, cx, py, cw, ph, 1.1, narrow=True)
    caption(s, 1.5, 5.35, 3.9, "Narrow: art above metadata; transport/volume scroll.")
    sidebar(s, ["Art 56 px in compact", "Metadata below art in narrow", "Volume label above slider", "Page scrolls vertically"], "Narrow art / metadata order")
    s.set_notes("Source: view.js adaptPages mediaRow/playerVolumeRow/outputVolumeRow vertical.")

    page_slide(d, "N14", "Alerts page", "Alerts — empty", "560×320 · centered empty-state card", "notifications",
               lambda c, x, y, w, h, z: draw_alerts(c, x, y, w, h, z, populated=False, indent=10),
               ["Count + Clear all header", "Weekday date subtitle", "Empty icon 32 px", "Title + one-line body centered"], "Empty card height / icon / copy",
               trigger="Clear all disabled when count=0. Shared emptyState() rhythm.", your_edits=["Count/date indented to card padding."])
    d.slides[-1].set_notes("Source: view.js buildNotifications/updateNotifications; widgets.js emptyState. Uploaded header indent preserved.")

    page_slide(d, "N15", "Alerts page", "Alerts — populated", "Notification cards keep native actions; list scrolls vertically", "notifications",
               lambda c, x, y, w, h, z: draw_alerts(c, x, y, w, h, z, populated=True, indent=10),
               ["Card radius 14 / padding 10", "List spacing 8 px", "Body color #C7C7CC", "Compact mode stacks actions"], "Card padding / action buttons / density",
               trigger="22 px source icon · title · dismiss · wrapped body · up to 2 native actions.", your_edits=["Count/date indented to card padding."])
    d.slides[-1].set_notes("Source: view.js updateNotifications; notifications.js. Uploaded header indent preserved.")

    s = d.slide("N51", "Alerts page", "Alerts — long body and scroll", "Fixed page height; vertical scroll preserves full content")
    c, cx, py, cw, ph, rect, foot = open_notch(s, "notifications", x=0.95, y=1.9, wpx=560, hpx=320, z=1.15, active="notifications")
    c.text(cx, py, cw * 0.55, px(22, 1.15), "3 notifications", size=fpx(14.7, 1.15), color=N_TEXT, bold=True, wrap=False)
    ui_button(c, cx + cw - px(100, 1.15), py, px(100, 1.15), px(28, 1.15), text="Clear all", ic="trash", kind="small", z=1.15)
    c.text(cx + px(10, 1.15), py + px(26, 1.15), cw, px(18, 1.15), "Friday, October 30", size=fpx(12.6, 1.15), color=N_MUTED, wrap=False)
    ly = py + px(52, 1.15)
    c.rect(cx, ly, cw, px(120, 1.15), fill=N_CARD, line=N_BORDER, radius=px(14, 1.15))
    c.text(cx + px(10, 1.15), ly + px(8, 1.15), cw - px(50, 1.15), px(20, 1.15), "Long notification title that ellipsizes", size=fpx(14.7, 1.15), color=N_TEXT, bold=True, wrap=False)
    c.text(cx + px(10, 1.15), ly + px(34, 1.15), cw - px(20, 1.15), px(76, 1.15), "Body wraps across multiple lines. The card grows to fit the text, and the list scrolls when the total height exceeds the page viewport. No body text is silently clipped.", size=fpx(13.3, 1.15), color="C7C7CC", valign="top", wrap=True)
    c.rect(cx + cw - px(6, 1.15), py + px(52, 1.15), px(4, 1.15), ph - px(52, 1.15), fill=N_CONTROL, radius=px(2, 1.15))
    c.rect(cx + cw - px(6, 1.15), py + px(60, 1.15), px(4, 1.15), px(60, 1.15), fill=N_MUTED, radius=px(2, 1.15))
    caption(s, 0.95, 6.1, 7.55, "Title ellipsizes; body wraps; list scrolls. Critical alerts still use native banners.")
    sidebar(s, ["Vertical scroll only", "Card grows with body", "Actions remain below body", "Clear all affects GNOME sources"], "Long-text density / scroll affordance")
    s.set_notes("Source: view.js updateNotifications wrapLabel; pageScrolls vscrollbar AUTOMATIC.")

    page_slide(d, "N16", "System page", "System — live stats", "640×260 · four stat cards + quick actions", "system",
               lambda c, x, y, w, h, z: draw_system(c, x, y, w, h, z, columns=4),
               ["4 columns ≥560 px content", "Card radius 14 / padding 12", "Value 1.45 em / 700", "Full action labels on System"], "Stat card order / value format / detail lines",
               trigger="Sampling only while System visible. Network uses default IPv4 interface; battery shows AC when absent.")
    d.slides[-1].set_notes("Source: view.js buildSystem/updateSystem; diagnostics.js; layout.js statColumns.")

    page_slide(d, "N17", "System page", "System — waiting / unavailable", "Initial sample and unavailable sensors use explicit placeholders", "system",
               lambda c, x, y, w, h, z: draw_system(c, x, y, w, h, z, columns=4, loading=True),
               ["No fake zero values", "Battery absent → AC", "Network absent → No interface", "Quick actions remain usable"], "Placeholder glyphs / unavailable copy",
               trigger="CPU/network use '…' while sampling; memory/battery use '—' when unavailable.")
    d.slides[-1].set_notes("Source: view.js updateSystem; diagnostics values.")

    s = d.slide("N52", "System page", "System — 2-column and 1-column", "statColumns(width): 4 ≥560 · 2 ≥320 · 1 <320")
    c, cx, py, cw, ph, rect, foot = open_notch(s, "system", x=0.55, y=2.0, wpx=480, hpx=260, z=0.95, active="system", header="active")
    draw_system_scroll(c, cx, py, cw, ph, 0.95, columns=2)
    caption(s, 0.55, 4.95, 4.6, "2 columns: stats fit; quick actions scroll into view.")
    c, cx, py, cw, ph, rect, foot = open_notch(s, "system", x=5.15, y=2.0, wpx=340, hpx=260, z=0.95, active="system", header="icons")
    draw_system_scroll(c, cx, py, cw, ph, 0.95, columns=1)
    caption(s, 5.15, 4.95, 3.25, "1 column: cards stack; page scrolls.")
    sidebar(s, ["Cards reparent into rows", "Quick actions stack in compact", "Sampling remains page-scoped", "No horizontal scroll"], "Column breakpoints / card density")
    s.set_notes("Source: layout.js statColumns; view.js layoutSystemCards/adaptPages.")

    s = d.slide("N18", "Quick actions", "Off · on · locked · hidden", "Shared by Home compact and System full labels")
    c = s.root
    y = 2.0
    for label, states, note in [("Off", (False, False, False), "Control fill #2C2C2E; text #E5E5EA"), ("On", (True, True, True), "Active fill #F5F5F7; text #000000"), ("Mixed", (False, True, True), "Each toggle reflects its own GSettings key"), ("Locked", (False, False, True), "Policy-locked keys are non-reactive; error explains lock")]:
        c.text(0.55, y, 1.2, 0.3, label, size=13, color=TEXT, bold=True, wrap=False)
        quick_actions(c, 1.9, y, 1.15, False, states)
        c.text(1.9, y + 0.45, 6.8, 0.25, note, size=10, color=MUTED, wrap=False)
        y += 1.05
    c.text(0.55, 6.3, 8.3, 0.3, "Hidden: Settings → Behavior → Quick actions removes the button from both pages.", size=10.5, color=MUTED, wrap=False)
    sidebar(s, ["Focus = inverted show-banners", "Night = night-light-enabled", "Dark = color-scheme prefer-dark", "Locked → '… is locked by system policy.'"], "Label length / active contrast / locked style")
    s.set_notes("Source: actions.js DEFINITIONS/toggle/writable; view.js syncActions/toggleAction.")

    page_slide(d, "N19", "Shelf page", "Shelf — ready, empty", "640×260 · helper connected · no references yet", "shelf",
               lambda c, x, y, w, h, z: draw_shelf(c, x, y, w, h, z, "ready", []),
               ["Header: Quick stash / Add / Clear", "Retry hidden when ready", "Empty state centered", "Dashed drop hint radius 14"], "Header buttons / empty copy / drop hint",
               trigger="Add files opens GTK file chooser. Clear disabled when empty.")
    d.slides[-1].set_notes("Source: view.js buildShelf/updateShelf; shelf.js; helper.py choose_files.")

    page_slide(d, "N20", "Shelf page", "Shelf — with files", "Rows store references only; originals are never moved or deleted", "shelf",
               lambda c, x, y, w, h, z: draw_shelf(c, x, y, w, h, z, "ready", ["Fedora-Workstation-44.iso", "design-review.pptx", "notes.md"]),
               ["Row radius 12 / padding 8×10", "Filename button expands", "Three icon actions at right", "Clear removes references only"], "Row height / filename truncation / icon order",
               trigger="Row actions: open · copy URI · reveal in Files · remove reference.")
    d.slides[-1].set_notes("Source: view.js updateShelf; shelf.js command; helper.py Open/Reveal/CopyUri/Remove.")

    s = d.slide("N21", "Shelf page", "Shelf — starting / offline", "Retry is visible whenever helper state is not ready")
    c, cx, py, cw, ph, rect, foot = open_notch(s, "shelf", y=1.9, hpx=260, z=1.18, active="shelf", status={"height": 46})
    draw_shelf(c, cx, py, cw, ph, 1.18, "offline", [])
    status_footer(c, foot[0], foot[1], foot[2], foot[3], "File shelf helper is offline. Press Retry to start it.", 1.18, error=True, retry=True)
    caption(s, 0.55, 5.55, 8.35, "Starting: Retry becomes 'Starting…' and is disabled. Offline: Retry re-activates the helper.")
    g = s.group("shelf-starting-mini")
    notch_surface(g, 0.55, 6.0, 3.9, 0.7, "open", 0.55)
    g.text(0.75, 6.12, 3.5, 0.2, "Starting file shelf helper…", size=9, color=N_MUTED, wrap=False)
    ui_button(g, 3.55, 6.12, 0.7, 0.25, text="Starting…", kind="small", z=0.5, disabled=True)
    sidebar(s, ["State: ready / starting / offline", "Status footer wraps + Retry", "Error color #FF9F0A", "Helper idle exit after 5 min is normal"], "Offline copy / retry placement / footer height")
    s.set_notes("Source: view.js retryShelf/updateShelf/showError; shelf.js start/state; helper.py idle exit.")

    s = d.slide("N22", "Shelf page", "External GTK drop target", "Drag files near the notch → temporary GTK4 surface positioned by Mutter")
    c = s.root
    panel_bar(c, 0.55, 1.95, 8.35, 0.4)
    g = s.group("gtk-drop")
    g.rect(1.0, 2.5, 7.45, 3.2, fill="14231D", line="8EE4C0", line_w=2, radius=0.23, dash=True)
    icon(g, "drop", 4.45, 2.95, 0.5, "E4FFF1")
    g.text(1.2, 3.6, 7.05, 0.4, "Drop files to stash", size=20, color="E4FFF1", bold=True, align="center", wrap=False)
    g.text(1.5, 4.15, 6.45, 0.6, "Only shelf references are saved. Originals stay where they are.", size=12, color="ACD5BE", align="center", valign="top")
    caption(s, 0.55, 5.95, 8.35, "Default target ≈600×232 at 640×320 open geometry. After drop: 'Files stashed. You can release the drag.'")
    sidebar(s, ["BG #14231D · text #E4FFF1", "Dashed 2 px #8EE4C0", "Radius 22 · title 20 px bold", "Detail #ACD5BE · icon 48 px"], "Target colors / copy / size")
    s.set_notes("Source: helper.py _load_gtk/_make_window/show_drop_target; view.js dragPosition/dropRectangle/placeHelper.")

    s = d.slide("N53", "Shelf page", "Drop result and file chooser", "Success feedback · invalid drop · GTK FileDialog")
    c = s.root
    c.rect(0.55, 1.95, 4.0, 1.5, fill="14231D", line="8EE4C0", line_w=1.5, radius=0.18, dash=True)
    c.text(0.75, 2.15, 3.6, 0.25, "Drop success", size=12, color="E4FFF1", bold=True, wrap=False)
    c.text(0.75, 2.55, 3.6, 0.6, "Files stashed. You can release the drag.", size=11, color="ACD5BE", valign="top")
    c.rect(0.55, 3.75, 4.0, 1.5, fill="2A1D12", line=ERROR, line_w=1.5, radius=0.18, dash=True)
    c.text(0.75, 3.95, 3.6, 0.25, "Invalid / failed drop", size=12, color=ERROR, bold=True, wrap=False)
    c.text(0.75, 4.35, 3.6, 0.6, "Helper shows the error text; target stays until drag leaves or retry.", size=11, color="F5D0A9", valign="top")
    draw_file_chooser(s, 4.85, 1.95, 4.05, 3.3)
    caption(s, 4.85, 5.45, 4.05, "Add files → GTK FileDialog / portal · accept label 'Stash'.")
    sidebar(s, ["Only file:// references accepted", "Originals never moved/deleted", "Chooser can select multiple files", "Helper exits after idle timeout"], "Drop feedback / chooser labels")
    s.set_notes("Source: helper.py _on_drop/choose_files; storage.py; view.js endDrag.")

    page_slide(d, "N23", "Clips page", "Clips — manual only / empty", "clipboard-enabled=false by default · explicit Capture now still works", "clipboard",
               lambda c, x, y, w, h, z: draw_clipboard(c, x, y, w, h, z, enabled=False, items=[]),
               ["Title: 'Clipboard history · manual only'", "Capture now / Clear history", "RAM-only privacy note", "Empty state explains opt-in"], "Privacy copy / empty state / button labels",
               trigger="Privacy note always visible. Clear history disabled until an item exists.")
    d.slides[-1].set_notes("Source: view.js buildClipboard/updateClipboard; clipboard.js; schema clipboard-enabled=false.")

    page_slide(d, "N24", "Clips page", "Clips — history on", "Bounded RAM-only history · click a row to restore it", "clipboard",
               lambda c, x, y, w, h, z: draw_clipboard(c, x, y, w, h, z, enabled=True, items=[("https://github.com/pavit12301611/notch", "14:30"), ("Design review notes: spacing, icons, empty states", "14:22"), ("sudo dnf install gnome-extensions-app", "13:58")]),
               ["Default limit 20; range 5–50", "64 KiB per item; RAM only", "Cleared on lock/logout/disable", "Password hints / mixed formats skipped"], "Row density / preview length / time alignment",
               trigger="Preview: flattened whitespace, max 160 chars. Time: locale HH:MM.")
    d.slides[-1].set_notes("Source: view.js updateClipboard; clipboard.js history; settings.py privacy description.")

    s = d.slide("N54", "Clips page", "Clips — enabled empty, long text, narrow", "Opt-in empty state · 160-char preview · stacked timestamp")
    c, cx, py, cw, ph, rect, foot = open_notch(s, "clipboard", x=0.55, y=1.95, wpx=460, hpx=300, z=0.9, active="clipboard", header="active")
    draw_clipboard(c, cx, py, cw, ph, 0.95, enabled=True, items=[])
    caption(s, 0.55, 5.0, 4.35, "Enabled but empty: 'Copy plain text to begin.'")
    c, cx, py, cw, ph, rect, foot = open_notch(s, "clipboard", x=5.15, y=1.95, wpx=340, hpx=300, z=0.9, active="clipboard", header="icons")
    draw_clipboard(c, cx, py, cw, ph, 0.9, enabled=True, items=[("Long captured text preview that ellipsizes at 160 characters…", "14:30"), ("Second clip", "14:22")], compact=True, narrow=True)
    caption(s, 5.15, 5.0, 3.25, "Narrow: stacked time; icon header.")
    sidebar(s, ["Manual capture works when off", "Automatic capture only when on", "Preview max 160 chars", "Lock clears/pauses history"], "Empty copy / narrow row layout")
    s.set_notes("Source: view.js updateClipboard; clipboard.js configure/capture; layout.js narrow.")

    s = d.slide("N25", "Shared chrome", "Status footer patterns", "Success / info / error · optional Retry · dismiss always available")
    c = s.root
    y = 1.95
    for text, err, retry, note in [("File URI copied.", False, False, "Transient success · 1800 ms"), ("Do Not Disturb on", False, False, "Quick action confirmation · 2200 ms"), ("File shelf helper could not start. Run notch-doctor.", True, True, "Helper error · Retry + dismiss"), ("Night Light is locked by system policy.", True, False, "Policy error · no retry")]:
        g = s.group(f"status-{y}")
        notch_surface(g, 0.55, y, 8.35, 0.95, "open", 0.8)
        status_footer(g, 0.85, y + 0.35, 7.75, 0.5, text, 0.95, error=err, retry=retry)
        c.text(0.55, y + 1.0, 8.35, 0.22, note, size=9.5, color=MUTED, wrap=False)
        y += 1.22
    sidebar(s, ["Info #FFD60A · error #FF9F0A", "Height = lines×18 + 10", "1–3 lines; chars/line ≈ width/7", "Footer reduces page height, never overlays"], "Footer padding / retry width / dismiss icon")
    s.set_notes("Source: layout.js statusHeight; view.js showStatus/showError/clearStatus.")

    s = d.slide("N26", "Adaptive layout", "Header modes", "Based on effective content width, not raw window width")
    c = s.root
    y = 1.95
    for mode, threshold, desc in [("labels", "≥760 px", "All six labels visible"), ("active", "420–759 px", "Only active tab keeps its label"), ("icons", "<420 px", "All tabs icon-only; no ellipsis")]:
        c.text(0.55, y, 1.5, 0.25, mode, size=13, color=ACCENT, bold=True, wrap=False)
        c.text(0.55, y + 0.3, 1.5, 0.22, threshold, size=10, color=MUTED, wrap=False, font=MONO)
        g = s.group(f"header-{mode}")
        w = 7.0 if mode == "labels" else (5.2 if mode == "active" else 3.4)
        g.rect(2.2, y, w, 0.62, fill=N_BG, radius=0.1)
        open_header(g, 2.35, y + 0.08, w - 0.3, "system", mode, 0.85)
        c.text(2.2, y + 0.7, w, 0.22, desc, size=10, color=MUTED, wrap=False)
        y += 1.45
    sidebar(s, ["Tab text never ellipsizes", "Settings + close always visible", "Active tab fill rgba(255,255,255,.16)", "Horizontal tab scroll disabled"], "Breakpoints / active label / icon padding")
    s.set_notes("Source: layout.js headerMode; view.js adaptPages; stylesheet.css notch-tab*.")

    s = d.slide("N27", "Adaptive layout", "Page reflow rules", "Compact <420 px · narrow <340 px effective content width")
    c = s.root
    rules = [("Home", "Compact stacks clock above centered actions; narrow hides Full player; art 56 px."), ("Music", "Narrow stacks art above metadata; volume rows stack; art 56 px in compact."), ("Alerts", "Compact Clear all becomes icon-only; notification actions stack vertically."), ("System", "Stats: 4 columns ≥560, 2 columns ≥320, 1 column below; actions stack in compact."), ("Shelf", "Compact Add/Clear/Retry become icon-only; rows keep three action icons."), ("Clips", "Compact header buttons icon-only; narrow rows stack timestamp below preview."), ("Text scale", "Geometry reserves space for GNOME text scaling; content is not double-scaled.")]
    y = 1.9
    for title, body in rules:
        c.rect(0.55, y, 8.35, 0.62, fill=PANEL_BG, line=STROKE, radius=0.08)
        c.text(0.75, y, 1.25, 0.62, title, size=11.5, color=ACCENT, bold=True, wrap=False)
        c.text(2.05, y, 6.65, 0.62, body, size=10.5, color=TEXT, valign="middle")
        y += 0.7
    sidebar(s, ["Breakpoints are logical effective px", "Vertical page scroll remains available", "No horizontal page scroll", "Minimum content 120×80"], "Add a custom breakpoint / reflow sketch")
    s.set_notes("Source: layout.js contentLayout/statColumns; view.js adaptPages.")

    s = d.slide("N55", "Adaptive layout", "Text scaling and monitor fit", "GNOME text-scaling-factor · notch-scale · monitor clamping")
    rows = [["100% text / 1.00 notch", "Default logical sizes; header mode from effective width."], ["125% text / 1.00 notch", "Geometry reserves larger text; effective width shrinks; labels may switch to icons."], ["150% text / 1.00 notch", "More page scrolling; no double scaling of content actor."], ["100% text / 1.50 notch", "Overall notch larger; monitor available size still clamps width/height."], ["Small monitor", "Open size clamps to monitor minus 12 px edge margins and placement top offset."]]
    native_table(s, 0.55, 1.95, 8.35, 3.6, ["Case", "Expected behavior"], rows, [2.6, 5.75], size=10.5)
    s.root.text(0.55, 5.85, 8.35, 0.6, "Acceptance: no clipped labels/buttons at 100%, 125%, 150% text scaling. Tabs never ellipsize; pages scroll vertically.", size=11, color=WARN, valign="top")
    sidebar(s, ["textScale clamped 0.8–2.0 in view", "notch-scale range 0.75–1.5", "Edge margin 12 logical px", "Primary monitor index −1"], "Scaling edge case to add")
    s.set_notes("Source: view.js readOptions; layout.js contentLayout/availableOpenSize; geometry.js notchGeometry.")

    s = d.slide("N28", "Motion", "Opening, closing and resizing", "Mac spring by default; smooth easing optional; reduced-motion respected")
    c = s.root
    y = 1.95
    for title, val, desc in [("Open", "response 0.42 / damping 0.8", "Content fades in 180 ms after 70 ms; closed layer fades 100 ms."), ("Close", "response 0.45 / damping 1.0", "Content fades 90 ms; closed layer returns after 60 ms / 140 ms."), ("Resize / page switch", "response 0.38 / damping 0.8", "Geometry springs x/y/w/h/radii; timeline stops when settled."), ("Smooth profile", "open 420 ms / close 450 ms", "easeOutCubic; durations user-configurable 0–1000 ms.")]:
        c.rect(0.55, y, 8.35, 0.95, fill=PANEL_BG, line=STROKE, radius=0.1)
        c.text(0.75, y + 0.12, 2.25, 0.25, title, size=13, color=TEXT, bold=True, wrap=False)
        c.text(3.05, y + 0.12, 5.65, 0.25, val, size=10.5, color=ACCENT, wrap=False, font=MONO)
        c.text(0.75, y + 0.45, 7.95, 0.4, desc, size=10.5, color=MUTED, valign="top")
        y += 1.1
    c.text(0.55, 6.45, 8.35, 0.3, "Animations off or GNOME Enable Animations off → immediate snap. Visualizer only while closed and visible.", size=10, color=MUTED, wrap=False)
    sidebar(s, ["Hover delay 250 ms", "Hover-exit close 120 ms", "Peek timeout 5000 ms", "Auto-hide 0 = disabled"], "Motion feel / durations / overshoot")
    s.set_notes("Source: spring.js SPRING_PROFILES; view.js animateTo/motionDuration/scheduleClose/syncVisualizer.")

    s = d.slide("N56", "Motion", "State transition storyboard", "Closed → open → page resize → peek → closed")
    frames = [("Closed", "closed", "home"), ("Open Home", "open", "home"), ("Music resize", "open", "media"), ("Peek", "peek", "notifications"), ("Closed", "closed", "home")]
    for i, (label, mode, tab) in enumerate(frames):
        x = 0.45 + i * 2.52
        motion_frame(s, x, 2.1, 2.35, 1.9, label, mode, tab)
        if i < 4:
            s.root.text(x + 2.28, 2.9, 0.3, 0.3, "→", size=16, color=ACCENT, align="center", wrap=False)
    c = s.root
    c.rect(0.55, 4.55, 12.23, 1.6, fill=PANEL_BG, line=STROKE, radius=0.12)
    c.text(0.8, 4.75, 11.7, 0.25, "Transition rules", size=12, color=ACCENT, bold=True, wrap=False)
    c.text(0.8, 5.15, 11.7, 0.8, "Hover opens after delay · click toggles immediately · Escape/outside click closes · page switch resizes geometry · notification may interrupt closed state with peek · external drag suppresses peek and shows GTK target.", size=11, color=TEXT, valign="top")
    sidebar(s, ["Geometry and opacity animate together", "Content fade delayed on open", "Reduced motion snaps", "No animation timers when idle"], "Transition to add/change")
    s.set_notes("Source: view.js setExpanded/selectTab/showPeek/dragPosition; spring.js.")

    s = d.slide("N29", "Behavior", "Visibility and lifecycle matrix", "When the notch hides, pauses or falls back")
    native_table(s, 0.55, 1.95, 8.35, 4.2, ["Context", "Notch", "Notifications", "Clipboard / sampling"],
                 [["Normal desktop", "Visible", "Routed to notch", "Clipboard opt-in; stats only on System"], ["Fullscreen (hide on)", "Hidden", "Native banner path", "Visualizer/stats paused"], ["Overview (hide on)", "Hidden", "Native banner path", "Visualizer/stats paused"], ["Lock / greeter", "Hidden", "Native path", "Clipboard paused/cleared"], ["External drag", "Drop target", "Peek suppressed", "Shelf helper active"], ["Helper idle 5 min", "Shelf offline", "Unchanged", "Helper exits; Retry restarts"]],
                 [2.0, 1.3, 2.0, 3.05], size=9.5)
    sidebar(s, ["hide-on-fullscreen=true", "hide-in-overview=true", "reserve-space=false (experimental)", "Panel clock restored when notch hides"], "Add/remove a visibility rule")
    s.set_notes("Source: view.js syncVisibility/connectEnvironment; notifications.js configure; helper.py idle exit.")

    s = d.slide("N30", "Settings", "Two native settings surfaces", "Same GSettings schema; inline extension prefs and companion app")
    c = s.root
    for x, title, body in [(0.55, "Extension preferences (inline)", "Opened from notch gear or GNOME Extensions.\nPages: Appearance · Behavior · Privacy.\nNo placeholder launcher page.\nMacBook preset applies the full default set."), (6.85, "Companion app (GTK4 / Adw)", "notch-settings · 760×780 default window.\nPages: Appearance · Behavior · Privacy & shelf · About.\nAdds file chooser, reference-list path and live status.\nAbout shows extension/helper health and source link.")]:
        c.rect(x, 1.95, 5.9, 4.6, fill=PANEL_BG, line=STROKE, radius=0.14)
        c.text(x + 0.25, 2.15, 5.4, 0.3, title, size=15, color=TEXT, bold=True, wrap=False)
        c.text(x + 0.25, 2.6, 5.4, 2.6, body, size=12, color=MUTED, valign="top")
    c.text(0.55, 6.7, 12.2, 0.25, "Next slides show each page. Rows are schematic; exact labels/defaults are listed for editing.", size=10.5, color=MUTED, wrap=False)
    s.set_notes("Source: extension/prefs.js; app/notch_app/settings.py; shared schema.")

    s = d.slide("N31", "Settings · Appearance", "Shape & position", "Logical pixels; clamped to selected monitor")
    settings_window(s, "Appearance", [{"title": "Shape & position", "description": "Dimensions use logical pixels and are clamped to the selected monitor.", "rows": [row("Placement", "combo", "MacBook top edge (default)", "MacBook mode replaces the centered top-bar clock"), row("MacBook preset", "button", "Apply", "Top-edge notch, idle clock, notification routing, smooth animation"), row("Monitor", "spin", -1, "−1 = primary monitor"), row("Collapsed base width", "spin", 185), row("Collapsed height", "spin", 32), row("Maximum page width", "spin", 640), row("Maximum page height", "spin", 320)]}], subtitle="Scroll for more geometry rows →", title="Notch Preferences", tabs=["Appearance", "Behavior", "Privacy"])
    sidebar(s, ["Placement: panel / top-center / floating", "Monitor −1…15", "Collapsed 150–520 × 24–64", "Expanded 480–1100 × 200–720"], "Row order / labels / grouping")
    s.set_notes("Source: prefs.js appearance_page; settings.py appearance_page.")

    s = d.slide("N32", "Settings · Appearance", "Geometry switches + finish", "Lower part of Appearance page")
    settings_window(s, "Appearance", [{"title": "Shape & position (continued)", "rows": [row("Match panel height when closed", "switch", True), row("Adaptive MacBook corners", "switch", True), row("Floating top margin", "spin", 12), row("Horizontal offset", "spin", 0)]}, {"title": "Finish", "rows": [row("Corner radius", "spin", 24), {"title": "Opacity", "kind": "slider", "value": 1.0, "label": "1.00", "subtitle": "The notch remains opaque enough for readable controls"}, {"title": "Notch scale", "kind": "slider", "value": 0.333, "label": "1.00"}]}], h=4.8, title="Notch Preferences", tabs=["Appearance", "Behavior", "Privacy"])
    s.root.text(0.55, 6.68, 8.35, 0.22, "Opacity 0.50–1.00 · Notch scale 0.75–1.50 · corner radius applies when adaptive corners are off.", size=8.5, color=MUTED, wrap=False)
    sidebar(s, ["Match panel height default on", "Adaptive corners default on", "Opacity 1.00; scale 1.00", "Corner radius 8–40"], "Finish controls / slider labels")
    s.set_notes("Source: prefs.js/settings.py Finish group. Notch scale slider midpoint shown schematically; actual default 1.00.")

    s = d.slide("N33", "Settings · Behavior", "MacBook notch + motion", "Top-edge behavior, notification routing and animation profile")
    settings_window(s, "Behavior", [{"title": "MacBook notch", "description": "The default is a top-edge notch, not a floating Dynamic Island.", "rows": [row("Replace top-bar date & time", "switch", True), row("Show time when idle", "switch", True), row("Show notifications in the notch", "switch", True), row("Compact notification peeks", "switch", True), row("Peek duration (ms)", "spin", 5000), row("Closed music wings", "switch", True), row("Motion profile", "combo", "Mac spring")]}], subtitle="Scroll for durations and opening/hiding groups →", title="Notch Preferences", tabs=["Appearance", "Behavior", "Privacy"])
    sidebar(s, ["Peek 2000–15000 ms, step 500", "Motion: Mac spring / Smooth easing", "Smooth open 420 ms; close 450 ms", "Durations 0–1000 ms, step 20"], "Switch order / motion labels")
    s.set_notes("Source: prefs.js/settings.py behavior_page MacBook group.")

    s = d.slide("N34", "Settings · Behavior", "Opening, hiding, actions, media", "Lower part of Behavior page")
    settings_window(s, "Behavior", [{"title": "Opening & hiding", "rows": [row("Open on hover", "switch", True), row("Hover delay (ms)", "spin", 250), row("Auto-hide delay (ms)", "spin", 0, "0 disables"), row("Smooth animations", "switch", True), row("Hide on fullscreen", "switch", True), row("Hide in Activities overview", "switch", True), row("Reserve collapsed space", "switch", False, "Experimental")]}], h=4.35, title="Notch Preferences", tabs=["Appearance", "Behavior", "Privacy"])
    s.root.rect(0.55, 6.3, 8.35, 0.45, fill=PANEL_BG, line=STROKE, radius=0.08)
    s.root.text(0.75, 6.3, 7.95, 0.45, "Quick actions: Do Not Disturb · Night Light · Dark Mode   |   Media: Download remote album art   |   Shortcut: Super + `", size=9.5, color=TEXT, valign="middle")
    sidebar(s, ["Hover 0–1500 ms", "Auto-hide 0–60000 ms", "All three quick actions default on", "Remote art default on"], "Group order / experimental badge")
    s.set_notes("Source: prefs.js/settings.py Opening & hiding, Quick actions, Media groups; schema toggle-notch.")

    s = d.slide("N57", "Settings · Privacy", "Inline extension Privacy page", "Clipboard controls + companion launcher")
    settings_window(s, "Privacy", [{"title": "Clipboard history", "description": "RAM-only, cleared on lock/disable. Keep off when handling secrets.", "rows": [row("Remember copied text", "switch", False), row("History size", "spin", 20), row("Clear clipboard history", "button", "Clear", "Does not erase the current system clipboard")]}, {"title": "File shelf", "description": "Stores local file references only. Originals are never moved or deleted.", "rows": [row("Companion app", "button", "Open", "File chooser, shelf tools and installation status")]}], h=4.8, title="Notch Preferences", tabs=["Appearance", "Behavior", "Privacy"])
    sidebar(s, ["Inline page title: Privacy", "No reference-path row inline", "Companion app launcher row", "Shared clipboard keys"], "Inline privacy wording / launcher")
    s.set_notes("Source: extension/prefs.js Privacy page.")

    s = d.slide("N35", "Settings · Privacy", "Companion Privacy & shelf", "Reference-list path + Choose files action")
    settings_window(s, "Privacy & shelf", [{"title": "Clipboard history", "description": "RAM-only, bounded, cleared on lock/disable. Keep off when handling secrets.", "rows": [row("Remember copied text", "switch", False), row("History size", "spin", 20, "Entries"), row("Clear clipboard history", "button", "Clear", "Does not erase the current system clipboard")]}, {"title": "File shelf", "description": "Local file references only. Originals are never moved or deleted.", "rows": [row("Reference list", "text", "~/.local/state/notch/shelf.json"), row("Add files", "button", "Choose files", "GTK file chooser / desktop portal")]}], h=4.8)
    sidebar(s, ["Clipboard off by default", "Limit 5–50, step 5", "64 KiB/item; 1 s capture poll", "Shelf stores references, not copies"], "Privacy wording / row order / warnings")
    s.set_notes("Source: settings.py privacy_page; helper/storage; schema defaults.")

    s = d.slide("N36", "Settings · About", "About + live status", "Companion app only")
    settings_window(s, "About", [{"title": "Notch 0.1.0", "description": "A native Shell notch with GTK4/Libadwaita companions. Inspired by Cyclop; no Cyclop code or assets are included.", "rows": [row("Shell extension", "text", "Running · 0 clips · 0 files"), row("File helper", "text", "Ready · version 0.1.0"), row("Refresh status", "button", "Check")]}, {"title": "Project", "rows": [row("Source & documentation", "button", "Open", "github.com/pavit12301611/notch")]}], h=4.7)
    sidebar(s, ["Extension status via GetState", "Helper status via Ping", "Offline errors shown inline", "Source link opens browser"], "Status copy / About layout")
    s.set_notes("Source: settings.py about_page/refresh_status.")

    s = d.slide("N58", "Settings", "Companion Appearance / Behavior", "Shared rows match inline prefs; window chrome differs")
    settings_window(s, "Appearance", [{"title": "Shared settings", "description": "Companion uses the same GSettings keys and row labels as inline prefs.", "rows": [row("Placement", "combo", "MacBook top edge (default)"), row("MacBook preset", "button", "Apply"), row("Collapsed base width", "spin", 185), row("Maximum page width", "spin", 640), row("Open on hover", "switch", True), row("Remember copied text", "switch", False)]}], h=4.8)
    sidebar(s, ["Window default 760×780", "Four tabs including About", "Behavior adds shortcut info row", "No duplicate schema keys"], "Companion-specific chrome")
    s.set_notes("Source: settings.py appearance_page/behavior_page; shared schema.")

    s = d.slide("N37", "Components", "Reusable control library", "Edit these master components; page mockups should stay consistent")
    c = s.root
    y = 1.95
    c.text(0.55, y, 2.0, 0.25, "BUTTONS", size=10, color=ACCENT, bold=True, wrap=False)
    ui_button(c, 0.55, y + 0.35, 1.2, 0.38, text="Button", kind="button", z=1.0)
    ui_button(c, 1.9, y + 0.35, 1.2, 0.38, text="Active", kind="action", z=1.0, active=True)
    ui_button(c, 3.25, y + 0.35, 1.2, 0.38, text="Disabled", kind="button", z=1.0, disabled=True)
    ui_button(c, 4.6, y + 0.35, 0.42, 0.38, ic="gear", kind="icon", z=1.0)
    ui_button(c, 5.15, y + 0.35, 0.55, 0.38, ic="play", kind="play", z=1.0)
    ui_button(c, 5.85, y + 0.35, 1.2, 0.38, text="Small", kind="small", z=1.0)
    y += 1.1
    c.text(0.55, y, 2.0, 0.25, "TAB / SLIDER", size=10, color=ACCENT, bold=True, wrap=False)
    ui_button(c, 0.55, y + 0.35, 1.3, 0.36, text="Home", ic="home", kind="tab", z=1.0, active=True)
    ui_button(c, 1.95, y + 0.35, 1.3, 0.36, text="Music", ic="music", kind="tab", z=1.0)
    slider(c, 3.6, y + 0.53, 3.0, 0.4, 1.0)
    y += 1.1
    c.text(0.55, y, 2.0, 0.25, "CARD / ROW / EMPTY", size=10, color=ACCENT, bold=True, wrap=False)
    c.rect(0.55, y + 0.35, 2.3, 1.0, fill=N_CARD, line=N_BORDER, radius=0.15)
    c.text(0.7, y + 0.45, 2.0, 0.2, "Card title", size=11, color=N_TEXT, bold=True, wrap=False)
    c.text(0.7, y + 0.7, 2.0, 0.5, "Body text wraps inside the card.", size=9.5, color=N_MUTED, valign="top")
    c.rect(3.05, y + 0.35, 3.0, 0.44, fill=N_CARD, radius=0.12)
    c.text(3.2, y + 0.35, 2.7, 0.44, "Row label", size=10.5, color=N_TEXT, wrap=False)
    empty_state(c, 6.25, y + 0.35, 2.5, 1.0, "folder", "Empty title", "Short explanation.", 0.7)
    sidebar(s, ["Button radius 10; play 18", "Tab active white 16%", "Card radius 14; row 12", "Slider bar 4 / handle 12"], "Component-level changes")
    s.set_notes("Source: stylesheet.css; widgets.js. Editable master components.")

    s = d.slide("N38", "Settings", "Inline prefs vs companion app", "Keep shared rows identical; companion-only rows are explicit")
    native_table(s, 0.55, 2.0, 12.23, 3.4, ["Page", "Inline extension prefs", "Companion app"],
                 [["Appearance", "Same geometry + finish rows", "Same rows"], ["Behavior", "Same behavior rows", "Same rows + shortcut info row"], ["Privacy", "Clipboard + Companion app launcher", "Privacy & shelf: reference path + Choose files"], ["About", "Not present", "Version, live status, refresh, source link"]],
                 [1.7, 5.0, 5.5], size=10.5)
    s.root.text(0.55, 5.8, 12.23, 0.5, "Rule: a shared GSettings key must not get two different labels/defaults across surfaces. Edit both columns if a change is intentional.", size=11, color=WARN, valign="top")
    s.set_notes("Source: extension/prefs.js vs app/notch_app/settings.py.")

    s = d.slide("N39", "Interactions", "Input and navigation map", "Every visible control has a defined action")
    c = s.root
    pairs = [("Click closed notch", "Toggle open / open Alerts from peek"), ("Hover 250 ms", "Open if open-on-hover; leave schedules close"), ("Super + `", "Toggle notch; configurable/disablable in schema"), ("Escape / outside click", "Collapse open notch"), ("Tab click", "Switch page; resize to page geometry"), ("Gear / close", "Open inline prefs / collapse"), ("Notification Open", "Activate source and open Alerts"), ("Shelf drag near notch", "Show GTK drop target; suppress peek"), ("Clip row click", "Restore text to system clipboard"), ("Quick action click", "Toggle GSettings; show status footer")]
    y = 1.9
    for a, b in pairs:
        c.rect(0.55, y, 8.35, 0.44, fill=PANEL_BG, line=STROKE, radius=0.06)
        c.text(0.75, y, 2.7, 0.44, a, size=10.5, color=ACCENT, bold=True, wrap=False)
        c.text(3.5, y, 5.2, 0.44, b, size=10.5, color=TEXT, wrap=False)
        y += 0.48
    sidebar(s, ["Keyboard focus rings use hover style", "Slider grabs postpone auto-hide", "Disabled controls stay visible", "No drag-out from shelf implemented"], "Add/change an interaction")
    s.set_notes("Source: view.js connectEnvironment/build; extension.js shortcut; shelf.js.")

    s = d.slide("N40", "Guardrails", "Performance, privacy and platform rules", "Design edits must not break these constraints")
    c = s.root
    cards = [("Performance", "Stats sample only on System page. Seek updates only on Music page. Visualizer only while closed/visible. Helper exits after 5 idle minutes."), ("Privacy", "Clipboard off by default, RAM-only, bounded, cleared on lock/logout/disable. Shelf stores references only. Remote art can be disabled."), ("Platform", "GNOME Wayland Shell extension for the notch; GTK4/Adw for settings/drop helper. No X11-only positioning, no floating-island default."), ("Accessibility", "Respect GNOME text scaling and Enable Animations. Symbolic icons inherit control color. Disabled/locked states remain understandable.")]
    for i, (title, body) in enumerate(cards):
        x, y = 0.55 + (i % 2) * 6.2, 1.95 + (i // 2) * 2.25
        c.rect(x, y, 5.95, 2.0, fill=PANEL_BG, line=STROKE, radius=0.14)
        c.text(x + 0.25, y + 0.2, 5.45, 0.3, title, size=15, color=ACCENT, bold=True, wrap=False)
        c.text(x + 0.25, y + 0.65, 5.45, 1.15, body, size=11.5, color=TEXT, valign="top")
    s.set_notes("Source: view.js setActive/syncVisualizer; helper.py idle; clipboard.js; docs/DESIGN_STUDY.md; README.")

    s = d.slide("N59", "Review worksheet", "Page-level decisions", "Fill one row per page; use exact values where possible")
    native_table(s, 0.55, 1.95, 12.23, 4.5, ["Page / slide", "KEEP / CHANGE / DISCUSS", "Exact edit", "Priority"],
                 [["Home N10–N11, N48", "", "", ""], ["Music N12–N13, N49–N50", "", "", ""], ["Alerts N14–N15, N51", "", "", ""], ["System N16–N18, N52", "", "", ""], ["Shelf N19–N22, N53", "", "", ""], ["Clips N23–N24, N54", "", "", ""]],
                 [2.6, 2.4, 5.6, 1.2], size=10.5)
    s.root.text(0.55, 6.65, 12.23, 0.25, "Priority: P0 must-fix · P1 should-fix · P2 nice-to-have.", size=10, color=MUTED, wrap=False)
    s.set_notes("Editable worksheet. Native PowerPoint table cells can be typed into directly.")

    s = d.slide("N60", "Review worksheet", "System-level decisions", "Tokens, geometry, settings, motion and guardrails")
    native_table(s, 0.55, 1.95, 12.23, 4.5, ["Area / slide", "KEEP / CHANGE / DISCUSS", "Exact edit", "Priority"],
                 [["Tokens N04 / N37", "", "", ""], ["Geometry N05 / N55", "", "", ""], ["Adaptive N26–N27", "", "", ""], ["Motion N28 / N56", "", "", ""], ["Settings N30–N36, N57–N58", "", "", ""], ["Guardrails N40", "", "", ""]],
                 [2.6, 2.4, 5.6, 1.2], size=10.5)
    s.root.text(0.55, 6.65, 12.23, 0.25, "If a change affects multiple pages, write the token/component change once and list affected slide IDs.", size=10, color=MUTED, wrap=False)
    s.set_notes("Editable worksheet for cross-cutting decisions.")

    s = d.slide("N41", "Handoff", "Decision log — fill this before returning", "One row per change; reference slide IDs so implementation is exact")
    native_table(s, 0.55, 1.95, 12.23, 4.5, ["Slide ID", "Element / state", "Requested change", "Priority", "Approved?"],
                 [["", "", "", "", ""] for _ in range(7)], [1.0, 2.4, 5.6, 1.2, 1.5], size=10.5)
    s.root.text(0.55, 6.65, 12.23, 0.25, "Priority: P0 must-fix · P1 should-fix · P2 nice-to-have. Duplicate this slide if you need more rows.", size=10, color=MUTED, wrap=False)
    s.set_notes("Editable decision log. Implementation should treat approved rows as the source of truth.")

    s = d.slide("N42", "Handoff", "Slide → source mapping", "After approval, edits are applied to these files")
    c = s.root
    mapping = [("N04, N37", "extension/stylesheet.css", "Colors, type, radii, spacing, control states"), ("N05–N08, N26–N28, N46–N56", "extension/lib/geometry.js · layout.js · surface.js · spring.js", "Dimensions, placement, adaptive rules, motion"), ("N09–N25, N39, N48–N54", "extension/lib/view.js · widgets.js", "Page structure, copy, controls, status footer"), ("N18", "extension/lib/actions.js", "Quick-action definitions and locked behavior"), ("N19–N22, N53", "extension/lib/shelf.js · app/notch_app/helper.py", "Shelf rows, drop target, file chooser"), ("N23–N24, N54", "extension/lib/clipboard.js", "History rules and restore"), ("N30–N36, N38, N57–N58", "extension/prefs.js · app/notch_app/settings.py", "Settings pages and rows"), ("N05, N31–N34, N55", "extension/schemas/*.gschema.xml", "Defaults, ranges, enums")]
    y = 1.9
    for ids, files, what in mapping:
        c.rect(0.55, y, 12.23, 0.56, fill=PANEL_BG, line=STROKE, radius=0.06)
        c.text(0.75, y, 2.35, 0.56, ids, size=9.5, color=ACCENT, bold=True, wrap=False)
        c.text(3.15, y, 5.45, 0.56, files, size=9, color=TEXT, wrap=False, font=MONO)
        c.text(8.75, y, 3.85, 0.56, what, size=9.5, color=MUTED, wrap=False)
        y += 0.6
    s.set_notes("Implementation map. No code changes are made by this deck itself.")

    s = d.slide("N43", "Handoff", "Implementation acceptance checklist", "Run after applying the edited deck")
    c = s.root
    checks = ["Every approved slide ID is implemented or explicitly marked deferred.", "No text is clipped at 100%, 125% and 150% text scaling.", "Closed / peek / open and all six pages match the edited mockups.", "Empty, loading, disabled, locked, offline and error states are present.", "Settings labels/defaults stay consistent across inline prefs and companion app.", "Quick actions, media controls, shelf actions and clipboard restore still work.", "Reduced motion, fullscreen/overview hide and lock privacy behavior still work.", "npm run check, Python tests, ruff, staged install/uninstall and notch-doctor pass."]
    y = 1.95
    for item in checks:
        c.rect(0.65, y + 0.05, 0.24, 0.24, fill=None, line=ACCENT, radius=0.04)
        c.text(1.05, y, 11.5, 0.5, item, size=12, color=TEXT, valign="middle")
        y += 0.58
    sidebar(s, ["Live GNOME validation happens on Fedora", "Deck is design truth, not a runtime screenshot", "Return .pptx with decision log filled"], "Add project-specific acceptance test")
    s.set_notes("Checklist mirrors existing CI and v3 standing requirement: no clipped text, working controls, adaptive UI.")

    s = d.slide("N61", "Appendix", "Uploaded install command", "Preserved verbatim from your main-branch PPT edit")
    c = s.root
    c.text(0.55, 1.95, 12.23, 0.35, "You added this command to the peek slide. It is preserved here so the notification mockup stays readable.", size=12, color=MUTED, wrap=False)
    c.rect(0.55, 2.55, 12.23, 1.3, fill="111118", line=STROKE, radius=0.12)
    c.text(0.8, 2.8, 11.73, 0.8, "curl -fsSL https://raw.githubusercontent.com/pavit12301611/notch/refs/heads/arena/a8df1ad3-notch/scripts/bootstrap.sh | bash", size=12, color=TEXT, font=MONO, valign="middle", wrap=True)
    c.text(0.55, 4.2, 12.23, 0.8, "Note: this is the historical branch command from the uploaded text. For the merged main branch, the equivalent raw URL uses refs/heads/main. Design review does not change installer behavior.", size=11, color=WARN, valign="top")
    s.set_notes("Verbatim user-added textbox from uploaded N08. Moved to appendix to preserve content without overlapping the peek mockup.")

    s = d.slide("N44", "Handoff", "Return the edited deck", "Next pass will follow the deck instead of guessing")
    c = s.root
    c.text(0.55, 2.1, 12.2, 0.5, "1. Edit shapes/text/colors directly on the page slides.", size=18, color=TEXT, bold=True, wrap=False)
    c.text(0.55, 2.75, 12.2, 0.5, "2. Fill REVIEW boxes, worksheets N59–N60 and decision log N41.", size=18, color=TEXT, bold=True, wrap=False)
    c.text(0.55, 3.4, 12.2, 0.5, "3. Save as .pptx and send it back.", size=18, color=TEXT, bold=True, wrap=False)
    c.text(0.55, 4.25, 12.2, 0.8, "Implementation will diff the edited deck against this baseline,\nmap changes through N42, and ask only where a deck edit is ambiguous.", size=14, color=MUTED, valign="top")
    chip(s, 0.55, 5.5, "BASELINE BRANCH: arena/71fe92f9-notch", ACCENT, 4.4)
    chip(s, 5.15, 5.5, "FILE: Notch-Design-Review.pptx", BLUE, 3.6)
    s.set_notes("End of deck. Keep slide IDs stable. New states may be added as Nxx-A.")

    return d


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--output", default="Notch-Design-Review.pptx")
    ap.add_argument("--preview", default=None, help="Optional directory for PNG previews/contact sheets")
    ap.add_argument("--warnings", default=None, help="Optional text file for layout warnings")
    args = ap.parse_args()
    deck = build_deck()
    out = Path(args.output)
    deck.save(out)
    if args.preview:
        deck.preview(Path(args.preview))
    if args.warnings:
        Path(args.warnings).parent.mkdir(parents=True, exist_ok=True)
        Path(args.warnings).write_text("\n".join(deck.warnings) + "\n")
    print(f"Saved {out} ({len(deck.slides)} slides, {out.stat().st_size / 1024:.1f} KiB)")
    if deck.warnings:
        print(f"Layout warnings: {len(deck.warnings)}")


if __name__ == "__main__":
    main()
